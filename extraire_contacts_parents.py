#!/usr/bin/env python3
"""Extraction et regroupement des contacts des parents des jeunes scouts SGDF.

Ce script lit en lecture seule le fichier d'export SGDF (`ExportIndividus.xls`),
extrait l'ensemble des représentants légaux (`1`, `2`, `3`) des jeunes du groupe,
les dédoublonne sur le couple `(Nom, Prénom)` afin de regrouper les fratries et
leurs unités courtes, enrichit les fiches avec les numéros d'adhérent et les rôles
adultes éventuels (ex. Parent Animateur Farfadet), trie les entrées par ordre
alphabétique français sur `(Nom, Prénom)` et génère le classeur Excel
`ContactsParents.xlsx` contenant les onglets `ExportIndividus` et `ContactsParents`.
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

FICHIER_ENTREE_DEFAUT = "ExportIndividus.xls"
FICHIER_SORTIE_DEFAUT = "ContactsParents.xlsx"

NOM_ONGLET_ORIGINE = "ExportIndividus"
NOM_ONGLET_PARENTS = "ContactsParents"

EN_TETES_CONTACTS_PARENTS: tuple[str, ...] = (
    "Nom",
    "Prénom",
    "Téléphone portable principal",
    "Téléphone portable secondaire",
    "Téléphone fixe (domicile)",
    "Adresse email",
    "Noms des enfants",
    "Unités des enfants",
    "Adresse physique",
    "Numéro d'adhérent",
    "Remarques",
)

LIBELLES_FONCTIONS_ADULTES: dict[tuple[str, str], str] = {
    ("271", "M."): "Parent Animateur Farfadet",
    ("271", "Mme"): "Parent Animateur Farfadet",
    ("270", "M."): "Responsable Farfadet",
    ("270", "Mme"): "Responsable Farfadet",
    ("300", "M."): "Responsable de groupe",
    ("300", "Mme"): "Responsable de groupe",
    ("301", "M."): "Responsable de groupe adjoint",
    ("301", "Mme"): "Responsable de groupe adjointe",
    ("307", "M."): "Secrétaire de groupe",
    ("307", "Mme"): "Secrétaire de groupe",
    ("309", "M."): "Trésorier de groupe",
    ("309", "Mme"): "Trésorière de groupe",
    ("316", "M."): "Cléophas de groupe",
    ("316", "Mme"): "Cléophas de groupe",
    ("330", "M."): "Chargé de mission du groupe",
    ("330", "Mme"): "Chargée de mission du groupe",
    ("181", "M."): "Membre du réseau Impeesa",
    ("181", "Mme"): "Membre du réseau Impeesa",
    ("380", "M."): "Coordonnateur du réseau Impeesa local",
    ("380", "Mme"): "Coordonnatrice du réseau Impeesa local",
    ("210", "M."): "Responsable d'unité Louveteaux-Jeannettes",
    ("210", "Mme"): "Responsable d'unité Louveteaux-Jeannettes",
    ("213", "M."): "Chef Louveteaux-Jeannettes",
    ("213", "Mme"): "Cheftaine Louveteaux-Jeannettes",
    ("220", "M."): "Responsable d'unité Scouts-Guides",
    ("220", "Mme"): "Responsable d'unité Scouts-Guides",
    ("223", "M."): "Chef Scouts-Guides",
    ("223", "Mme"): "Cheftaine Scouts-Guides",
    ("230", "M."): "Responsable d'unité Pionniers-Caravelles",
    ("230", "Mme"): "Responsable d'unité Pionniers-Caravelles",
    ("233", "M."): "Chef Pionniers-Caravelles",
    ("233", "Mme"): "Cheftaine Pionniers-Caravelles",
}


class ErreurExtractionContacts(Exception):
    """Erreur métier levée lors de la lecture ou de la validation de l'export SGDF."""


@dataclass(frozen=True)
class EnfantRattache:
    """Représente un jeune scout rattaché à un parent."""

    prenom: str
    nom: str
    unite_courte: str
    statut_inscription: str

    def formater_nom_complet(self) -> str:
        """Retourne le nom complet de l'enfant sous la forme 'Prénom NOM'."""
        return " ".join(part for part in (self.prenom, self.nom) if part)


@dataclass
class ProfilAdulteGroupe:
    """Représente une inscription adulte dans le groupe SGDF pour enrichir un parent."""

    nom: str
    prenom: str
    numero_adherent: str = ""
    roles_et_delegations: list[str] = field(default_factory=list)
    telephone_portable_1: str = ""
    telephone_portable_2: str = ""
    telephone_domicile: str = ""
    courriel: str = ""
    adresse_physique: str = ""


@dataclass
class ContactParent:
    """Représente la fiche consolidée et dédoublonnée d'un parent de jeune(s) scout(s)."""

    nom: str
    prenom: str
    telephone_portable_principal: str = ""
    telephone_portable_secondaire: str = ""
    telephone_fixe: str = ""
    adresse_email: str = ""
    enfants: list[EnfantRattache] = field(default_factory=list)
    unites: list[str] = field(default_factory=list)
    adresse_physique: str = ""
    numero_adherent: str = ""
    remarques_liste: list[str] = field(default_factory=list)

    def ajouter_enfant(self, enfant: EnfantRattache) -> None:
        """Ajoute un enfant et son unité sans doublon à la fiche du parent."""
        if enfant not in self.enfants:
            self.enfants.append(enfant)
        if enfant.unite_courte and enfant.unite_courte not in self.unites:
            self.unites.append(enfant.unite_courte)
        if enfant.statut_inscription and enfant.statut_inscription != "Inscrit(e)":
            mention_statut = f"Enfant {enfant.prenom} : {enfant.statut_inscription}"
            self.ajouter_remarque(mention_statut)

    def ajouter_remarque(self, remarque: str) -> None:
        """Ajoute une remarque non vide sans créer de doublon."""
        texte_propre = nettoyer_espaces(remarque)
        if texte_propre and texte_propre not in self.remarques_liste:
            self.remarques_liste.append(texte_propre)

    def fusionner_coordonnees(
        self,
        *,
        portable_principal: str = "",
        portable_secondaire: str = "",
        fixe: str = "",
        email: str = "",
        adresse: str = "",
        numero_adherent: str = "",
    ) -> None:
        """Conserve la première valeur non vide pour chaque coordonnée du parent."""
        if not self.telephone_portable_principal and portable_principal:
            self.telephone_portable_principal = portable_principal
        elif (
            not self.telephone_portable_secondaire
            and portable_principal
            and portable_principal != self.telephone_portable_principal
        ):
            self.telephone_portable_secondaire = portable_principal

        if (
            not self.telephone_portable_secondaire
            and portable_secondaire
            and portable_secondaire != self.telephone_portable_principal
        ):
            self.telephone_portable_secondaire = portable_secondaire

        if not self.telephone_fixe and fixe:
            self.telephone_fixe = fixe
        if not self.adresse_email and email:
            self.adresse_email = email
        if not self.adresse_physique and adresse:
            self.adresse_physique = adresse
        if not self.numero_adherent and numero_adherent:
            self.numero_adherent = numero_adherent

        self.reorganiser_telephones()

    def reorganiser_telephones(self) -> None:
        """Déplace un portable valide (+336/+337) vers portable 1 s'il est vide (priorité au fixe)."""
        if self.telephone_portable_principal:
            return

        if est_numero_portable_valide(self.telephone_fixe):
            self.telephone_portable_principal = self.telephone_fixe
            self.telephone_fixe = ""
            if self.telephone_portable_secondaire == self.telephone_portable_principal:
                self.telephone_portable_secondaire = ""
        elif est_numero_portable_valide(self.telephone_portable_secondaire):
            self.telephone_portable_principal = self.telephone_portable_secondaire
            self.telephone_portable_secondaire = ""

    @property
    def noms_enfants(self) -> str:
        """Retourne la liste dédoublonnée des enfants ('Prénom NOM') séparée par ', '."""
        libelles: list[str] = []
        for enfant in self.enfants:
            libelle = enfant.formater_nom_complet()
            if libelle and libelle not in libelles:
                libelles.append(libelle)
        return ", ".join(libelles)

    @property
    def unites_enfants(self) -> str:
        """Retourne les unités courtes des enfants dans le même ordre que les enfants eux-mêmes."""
        return ", ".join(
            enfant.unite_courte for enfant in self.enfants if enfant.unite_courte
        )

    @property
    def remarques(self) -> str:
        """Retourne les remarques consolidées séparées par ' ; '."""
        return " ; ".join(self.remarques_liste)

    def vers_ligne_excel(self) -> list[str]:
        """Convertit la fiche du parent en liste des 11 valeurs de colonnes attendues."""
        return [
            self.nom,
            self.prenom,
            self.telephone_portable_principal,
            self.telephone_portable_secondaire,
            self.telephone_fixe,
            self.adresse_email,
            self.noms_enfants,
            self.unites_enfants,
            self.adresse_physique,
            self.numero_adherent,
            self.remarques,
        ]


def nettoyer_espaces(valeur: str) -> str:
    """Supprime les espaces superflus en début, fin et milieu de chaîne."""
    return re.sub(r"\s+", " ", valeur.strip())


def normaliser_nom(nom_brut: str) -> str:
    """Normalise un nom de famille intégralement en MAJUSCULES."""
    return nettoyer_espaces(nom_brut).upper()


def normaliser_prenom(prenom_brut: str) -> str:
    """Normalise un prénom avec la première lettre de chaque partie en majuscule et le reste en minuscules."""
    texte = nettoyer_espaces(prenom_brut).lower()
    if not texte:
        return ""
    return re.sub(
        r"(^|[\s\-'])([^\s\-'])",
        lambda correspondance: correspondance.group(1) + correspondance.group(2).upper(),
        texte,
    )


def normaliser_telephone(numero_brut: str) -> str:
    """Normalise un numéro de téléphone au format international compact sans espaces (ex. +33687966975)."""
    texte = numero_brut.strip()
    if not texte:
        return ""

    # Supprimer l'indication facultative (0) parfois présente après l'indicatif (+33(0)6...)
    texte = re.sub(r"\(\s*0\s*\)", "", texte)

    a_prefixe_plus = texte.startswith("+")
    chiffres = re.sub(r"\D", "", texte)
    if not chiffres:
        return ""

    if a_prefixe_plus:
        numero_normalise = f"+{chiffres}"
    elif chiffres.startswith("00") and len(chiffres) > 2:
        numero_normalise = f"+{chiffres[2:]}"
    elif chiffres.startswith("0") and len(chiffres) > 1:
        numero_normalise = f"+33{chiffres[1:]}"
    elif len(chiffres) == 9 and chiffres[0] in "123456789":
        numero_normalise = f"+33{chiffres}"
    elif chiffres.startswith("33") and len(chiffres) == 11:
        numero_normalise = f"+{chiffres}"
    else:
        numero_normalise = f"+{chiffres}"

    # Corriger le cas +3306... (13 caractères avec un 0 superflu après +33)
    if numero_normalise.startswith("+330") and len(numero_normalise) == 13:
        numero_normalise = f"+33{numero_normalise[4:]}"

    return numero_normalise


def est_numero_portable_valide(numero: str) -> bool:
    """Indique si le numéro normalisé est un téléphone portable français valide (+336 ou +337)."""
    return numero.startswith(("+336", "+337"))


def sans_accents(texte: str) -> str:
    """Retourne une chaîne dépourvue de marques diacritiques pour le tri alphabétique français."""
    decomposition = unicodedata.normalize("NFD", texte)
    return "".join(car for car in decomposition if unicodedata.category(car) != "Mn")


def cle_tri_parent(contact: ContactParent) -> tuple[str, str, str, str]:
    """Clé de tri alphabétique français sur le Nom puis sur le Prénom."""
    return (
        sans_accents(contact.nom).upper(),
        contact.nom,
        sans_accents(contact.prenom).lower(),
        contact.prenom,
    )


def normaliser_unite_courte(nom_structure: str, code_fonction: str = "") -> str:
    """Convertit un nom de structure ou code fonction SGDF en nom d'unité court normalisé."""
    code = nettoyer_espaces(code_fonction)
    structure_maj = sans_accents(nettoyer_espaces(nom_structure)).upper()

    if code in ("170", "270", "271") or "FARFADET" in structure_maj:
        return "Farfadets"
    if code in ("110", "210", "213") or "LOUVETEAU" in structure_maj or "JEANNETTE" in structure_maj:
        return "Louveteaux-Jeannettes"
    if code in ("120", "220", "223") or "SCOUTS GUIDES" in structure_maj or "SCOUT" in structure_maj:
        return "Scouts-Guides"
    if code in ("130", "230", "233") or "PIONNIER" in structure_maj or "CARAVELLE" in structure_maj:
        return "Pionniers-Caravelles"
    if code in ("140", "240", "243") or "COMPAGNON" in structure_maj:
        return "Compagnons"
    return nettoyer_espaces(nom_structure)


def normaliser_role_adulte(
    civilite: str,
    code_fonction: str,
    nom_masculin: str,
    nom_feminin: str,
) -> str:
    """Retourne le libellé lisible en français du rôle adulte au sein du groupe SGDF."""
    code = nettoyer_espaces(code_fonction)
    civ = nettoyer_espaces(civilite)
    cle = (code, "Mme" if civ == "Mme" else "M.")
    if cle in LIBELLES_FONCTIONS_ADULTES:
        return LIBELLES_FONCTIONS_ADULTES[cle]

    brut = nom_feminin if (civ == "Mme" and nom_feminin.strip()) else nom_masculin
    texte = nettoyer_espaces(brut)
    if not texte:
        return ""
    return texte[0].upper() + texte[1:].lower()


def extraire_roles_delegations(texte_delegations: str, civilite: str = "M.") -> list[str]:
    """Extrait les rôles adultes mentionnés dans la colonne `Inscription.Delegations`."""
    texte = nettoyer_espaces(texte_delegations)
    if not texte:
        return []

    roles: list[str] = []
    for segment in texte.split(";"):
        morceau = segment.strip()
        if not morceau:
            continue
        correspondance = re.match(r"^(\d+)\s*-\s*([^(]+)", morceau)
        if correspondance:
            code_del = correspondance.group(1).strip()
            libelle_brut = correspondance.group(2).strip()
            role = normaliser_role_adulte(civilite, code_del, libelle_brut, libelle_brut)
            if role and role not in roles:
                roles.append(role)
        else:
            roles.append(morceau)
    return roles


def assembler_adresse_physique(
    ligne1: str,
    ligne2: str,
    ligne3: str,
    code_postal: str,
    municipalite: str,
    pays: str = "",
) -> str:
    """Assemble proprement les composantes d'une adresse postale."""
    elements: list[str] = []
    for ligne in (ligne1, ligne2, ligne3):
        propre = nettoyer_espaces(ligne)
        if propre:
            elements.append(propre)

    cp_propre = nettoyer_espaces(code_postal)
    ville_propre = nettoyer_espaces(municipalite).upper()
    bloc_ville = " ".join(part for part in (cp_propre, ville_propre) if part)
    if bloc_ville:
        elements.append(bloc_ville)

    pays_propre = nettoyer_espaces(pays).upper()
    if pays_propre and pays_propre != "FRANCE":
        elements.append(pays_propre)

    return ", ".join(elements)


def lire_tableau_export_individus(chemin_entree: Path) -> tuple[list[str], list[list[str]]]:
    """Lit en lecture seule le fichier HTML/XLS `ExportIndividus.xls` et retourne `(en_tetes, lignes)`."""
    if not chemin_entree.is_file():
        raise ErreurExtractionContacts(
            f"Le fichier d'entrée '{chemin_entree}' est introuvable."
        )

    try:
        contenu_octets = chemin_entree.read_bytes()
    except OSError as erreur:
        raise ErreurExtractionContacts(
            f"Impossible de lire le fichier '{chemin_entree}' : {erreur}"
        ) from erreur

    document = BeautifulSoup(contenu_octets, "lxml")
    lignes_tr = document.find_all("tr")
    if not lignes_tr:
        raise ErreurExtractionContacts(
            f"Aucun tableau HTML (`<tr>`) trouvé dans le fichier '{chemin_entree}'."
        )

    en_tetes = [
        nettoyer_espaces(cellule.get_text())
        for cellule in lignes_tr[0].find_all(["th", "td"])
    ]
    if not en_tetes or "Individu.Nom" not in en_tetes:
        raise ErreurExtractionContacts(
            f"En-têtes SGDF attendus (`Individu.Nom`) absents dans '{chemin_entree}'."
        )

    lignes_donnees: list[list[str]] = []
    nombre_colonnes = len(en_tetes)
    for ligne_tr in lignes_tr[1:]:
        cellules = [
            nettoyer_espaces(cellule.get_text())
            for cellule in ligne_tr.find_all("td")
        ]
        if not any(cellules):
            continue
        if len(cellules) < nombre_colonnes:
            cellules.extend([""] * (nombre_colonnes - len(cellules)))
        lignes_donnees.append(cellules[:nombre_colonnes])

    return en_tetes, lignes_donnees


def construire_index_adultes(
    en_tetes: list[str],
    lignes: list[list[str]],
) -> dict[tuple[str, str], ProfilAdulteGroupe]:
    """Construit un index des inscrits adultes (`Fonction.CategorieMembre != '0'`) par `(Nom, Prénom)`."""
    indice_col = {nom_col: idx for idx, nom_col in enumerate(en_tetes)}

    def val(ligne: list[str], colonne: str) -> str:
        idx = indice_col.get(colonne)
        return ligne[idx] if idx is not None and idx < len(ligne) else ""

    index_adultes: dict[tuple[str, str], ProfilAdulteGroupe] = {}
    for ligne in lignes:
        categorie = val(ligne, "Fonction.CategorieMembre")
        if categorie == "0":
            continue

        nom = normaliser_nom(val(ligne, "Individu.Nom"))
        prenom = normaliser_prenom(val(ligne, "Individu.Prenom"))
        if not nom and not prenom:
            continue

        cle = (nom, prenom)
        profil = index_adultes.get(cle)
        if profil is None:
            profil = ProfilAdulteGroupe(nom=nom, prenom=prenom)
            index_adultes[cle] = profil

        code_adherent = val(ligne, "Individu.CodeAdherent")
        if not profil.numero_adherent and code_adherent:
            profil.numero_adherent = code_adherent

        civilite = val(ligne, "IndividuCivilite.NomCourt")
        role_principal = normaliser_role_adulte(
            civilite=civilite,
            code_fonction=val(ligne, "Fonction.Code"),
            nom_masculin=val(ligne, "Fonction.NomMasculin"),
            nom_feminin=val(ligne, "Fonction.NomFeminin"),
        )
        if role_principal and role_principal not in profil.roles_et_delegations:
            profil.roles_et_delegations.append(role_principal)

        for role_del in extraire_roles_delegations(val(ligne, "Inscription.Delegations"), civilite):
            if role_del not in profil.roles_et_delegations:
                profil.roles_et_delegations.append(role_del)

        if not profil.telephone_portable_1:
            profil.telephone_portable_1 = normaliser_telephone(
                val(ligne, "Individu.TelephonePortable1")
            )
        if not profil.telephone_portable_2:
            profil.telephone_portable_2 = normaliser_telephone(
                val(ligne, "Individu.TelephonePortable2")
            )
        if not profil.telephone_domicile:
            profil.telephone_domicile = normaliser_telephone(
                val(ligne, "Individu.TelephoneDomicile")
            )
        if not profil.courriel:
            profil.courriel = (
                val(ligne, "Individu.CourrielPersonnel")
                or val(ligne, "Individu.CourrielDédiéSGDF")
            )
        if not profil.adresse_physique:
            profil.adresse_physique = assembler_adresse_physique(
                val(ligne, "Individu.Adresse.Ligne1"),
                val(ligne, "Individu.Adresse.Ligne2"),
                val(ligne, "Individu.Adresse.Ligne3"),
                val(ligne, "Individu.Adresse.CodePostal"),
                val(ligne, "Individu.Adresse.Municipalite"),
                val(ligne, "Individu.Adresse.Pays"),
            )

    return index_adultes


def extraire_et_regrouper_parents(
    en_tetes: list[str],
    lignes: list[list[str]],
) -> list[ContactParent]:
    """Extrait, dédoublonne, enrichit et trie par ordre alphabétique tous les parents de jeunes."""
    indice_col = {nom_col: idx for idx, nom_col in enumerate(en_tetes)}

    def val(ligne: list[str], colonne: str) -> str:
        idx = indice_col.get(colonne)
        return ligne[idx] if idx is not None and idx < len(ligne) else ""

    index_adultes = construire_index_adultes(en_tetes, lignes)
    parents_par_cle: dict[tuple[str, str], ContactParent] = {}

    for ligne in lignes:
        if val(ligne, "Fonction.CategorieMembre") != "0":
            continue

        enfant = EnfantRattache(
            prenom=normaliser_prenom(val(ligne, "Individu.Prenom")),
            nom=normaliser_nom(val(ligne, "Individu.Nom")),
            unite_courte=normaliser_unite_courte(
                val(ligne, "Structure.Nom"),
                val(ligne, "Fonction.Code"),
            ),
            statut_inscription=val(ligne, "Inscriptions.Type") or "Inscrit(e)",
        )

        for prefixe in ("RepresentantLegal1", "RepresentantLegal2", "RepresentantLegal3"):
            nom_parent = normaliser_nom(val(ligne, f"{prefixe}.Nom"))
            prenom_parent = normaliser_prenom(val(ligne, f"{prefixe}.Prenom"))
            if not nom_parent and not prenom_parent:
                continue

            cle_parent = (nom_parent, prenom_parent)
            contact = parents_par_cle.get(cle_parent)
            if contact is None:
                contact = ContactParent(nom=nom_parent, prenom=prenom_parent)
                parents_par_cle[cle_parent] = contact

            courriel_parent = (
                val(ligne, f"{prefixe}.CourrielPersonnel")
                or val(ligne, f"{prefixe}.CourrielDédiéSGDF")
            )
            adresse_parent = assembler_adresse_physique(
                val(ligne, f"{prefixe}.Adresse.Ligne1"),
                val(ligne, f"{prefixe}.Adresse.Ligne2"),
                val(ligne, f"{prefixe}.Adresse.Ligne3"),
                val(ligne, f"{prefixe}.Adresse.CodePostal"),
                val(ligne, f"{prefixe}.Adresse.Municipalite"),
                val(ligne, f"{prefixe}.Adresse.Pays"),
            )

            contact.fusionner_coordonnees(
                portable_principal=normaliser_telephone(
                    val(ligne, f"{prefixe}.TelephonePortable1")
                ),
                portable_secondaire=normaliser_telephone(
                    val(ligne, f"{prefixe}.TelephonePortable2")
                ),
                fixe=normaliser_telephone(
                    val(ligne, f"{prefixe}.TelephoneDomicile")
                ),
                email=courriel_parent,
                adresse=adresse_parent,
                numero_adherent=val(ligne, f"{prefixe}.CodeAdherent"),
            )
            contact.ajouter_enfant(enfant)

    # Enrichissement avec l'index des profils adultes inscrits dans le groupe
    for cle_parent, contact in parents_par_cle.items():
        profil_adulte = index_adultes.get(cle_parent)
        if profil_adulte is None:
            continue

        contact.fusionner_coordonnees(
            portable_principal=profil_adulte.telephone_portable_1,
            portable_secondaire=profil_adulte.telephone_portable_2,
            fixe=profil_adulte.telephone_domicile,
            email=profil_adulte.courriel,
            adresse=profil_adulte.adresse_physique,
            numero_adherent=profil_adulte.numero_adherent,
        )

        # Insérer les rôles adultes en tête des remarques (avant les statuts d'inscription des enfants)
        remarques_existantes = list(contact.remarques_liste)
        contact.remarques_liste.clear()
        for role in profil_adulte.roles_et_delegations:
            contact.ajouter_remarque(role)
        for remarque_enfant in remarques_existantes:
            contact.ajouter_remarque(remarque_enfant)

    return sorted(parents_par_cle.values(), key=cle_tri_parent)


def ajuster_largeur_colonnes(feuille) -> None:
    """Ajuste automatiquement la largeur des colonnes d'une feuille Excel pour une lecture confortable."""
    for indice_col, colonne in enumerate(feuille.columns, start=1):
        longueur_max = 0
        for cellule in colonne:
            valeur_texte = str(cellule.value or "")
            longueur_max = max(longueur_max, len(valeur_texte))
        lettre_colonne = get_column_letter(indice_col)
        feuille.column_dimensions[lettre_colonne].width = min(max(longueur_max + 3, 12), 55)


def ecrire_classeur_contacts_parents(
    chemin_sortie: Path,
    en_tetes_origine: list[str],
    lignes_origine: list[list[str]],
    contacts_parents: list[ContactParent],
) -> None:
    """Génère le classeur Excel `.xlsx` contenant les onglets `ExportIndividus` et `ContactsParents`."""
    chemin_sortie.parent.mkdir(parents=True, exist_ok=True)

    classeur = Workbook()
    police_entete = Font(bold=True)
    alignement_entete = Alignment(vertical="center")

    # Onglet 1 : ExportIndividus (copie fidèle des données d'origine au format texte)
    feuille_origine = classeur.active
    feuille_origine.title = NOM_ONGLET_ORIGINE
    feuille_origine.append(en_tetes_origine)
    for cellule in feuille_origine[1]:
        cellule.font = police_entete
        cellule.alignment = alignement_entete
        cellule.number_format = "@"

    for ligne_donnees in lignes_origine:
        feuille_origine.append(ligne_donnees)
        for cellule in feuille_origine[feuille_origine.max_row]:
            cellule.number_format = "@"
    feuille_origine.freeze_panes = "A2"

    # Onglet 2 : ContactsParents (liste dédoublonnée et triée des parents)
    feuille_parents = classeur.create_sheet(title=NOM_ONGLET_PARENTS)
    feuille_parents.append(list(EN_TETES_CONTACTS_PARENTS))
    for cellule in feuille_parents[1]:
        cellule.font = police_entete
        cellule.alignment = alignement_entete
        cellule.number_format = "@"

    for contact in contacts_parents:
        feuille_parents.append(contact.vers_ligne_excel())
        for cellule in feuille_parents[feuille_parents.max_row]:
            cellule.number_format = "@"
    feuille_parents.freeze_panes = "A2"
    ajuster_largeur_colonnes(feuille_parents)

    classeur.save(chemin_sortie)


def executer_extraction(chemin_entree: Path, chemin_sortie: Path) -> list[ContactParent]:
    """Exécute la chaîne complète : lecture seule, extraction, dédoublonnage, tri et écriture Excel."""
    en_tetes, lignes = lire_tableau_export_individus(chemin_entree)
    contacts_parents = extraire_et_regrouper_parents(en_tetes, lignes)
    ecrire_classeur_contacts_parents(
        chemin_sortie=chemin_sortie,
        en_tetes_origine=en_tetes,
        lignes_origine=lignes,
        contacts_parents=contacts_parents,
    )
    return contacts_parents


def analyser_arguments(arguments: list[str] | None = None) -> argparse.Namespace:
    """Analyse les arguments en ligne de commande (intégralement en français)."""
    analyseur = argparse.ArgumentParser(
        prog="extraire_contacts_parents.py",
        description=(
            "Extrait, dédoublonne et trie les contacts des parents des jeunes scouts SGDF "
            "depuis l'export intranet vers l'onglet 'ContactsParents'."
        ),
    )
    analyseur.add_argument(
        "-e",
        "--entree",
        type=Path,
        default=Path(FICHIER_ENTREE_DEFAUT),
        help=f"Chemin du fichier d'export SGDF en entrée (défaut : {FICHIER_ENTREE_DEFAUT}).",
    )
    analyseur.add_argument(
        "-s",
        "--sortie",
        type=Path,
        default=Path(FICHIER_SORTIE_DEFAUT),
        help=f"Chemin du classeur Excel généré en sortie (défaut : {FICHIER_SORTIE_DEFAUT}).",
    )
    return analyseur.parse_args(arguments)


def principale(arguments: list[str] | None = None) -> int:
    """Point d'entrée principal du script en ligne de commande."""
    options = analyser_arguments(arguments)
    try:
        contacts = executer_extraction(
            chemin_entree=options.entree,
            chemin_sortie=options.sortie,
        )
    except ErreurExtractionContacts as erreur:
        print(f"Erreur : {erreur}", file=sys.stderr)
        return 1

    print(
        f"Extraction terminée avec succès : {len(contacts)} parents uniques enregistrés "
        f"dans l'onglet '{NOM_ONGLET_PARENTS}' du fichier '{options.sortie}'."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(principale())
