"""Tests unitaires et d'intégration pour `extract_parent_contacts.py`.

Vérifie :
1. La normalisation des noms et prénoms (casse titre et composés) et unités courtes.
2. Le dédoublonnage des parents sur `(Nom, Prénom)` et le regroupement des fratries et unités.
3. L'enrichissement du numéro d'adhérent et des remarques (rôles adultes comme Parent Animateur
   Farfadet, délégations et statuts d'enfants `Pré-inscrit(e)` / `Invité(e)`).
4. Le tri alphabétique français (gestion des caractères accentués) sur `(Nom, Prénom)`.
5. La création du classeur `ContactsParents.xlsx` avec les onglets `ExportIndividus` et
   `ContactsParents`, ainsi que la lecture seule stricte de `ExportIndividus.xls`.
"""

from __future__ import annotations

import hashlib
import sys
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from extract_parent_contacts import (
    PARENT_CONTACTS_HEADERS,
    PARENTS_SHEET_NAME,
    SOURCE_SHEET_NAME,
    extract_and_group_parents,
    join_street_lines,
    main,
    normalize_address_part,
    normalize_first_name,
    normalize_last_name,
    normalize_phone,
    normalize_short_unit,
    run_extraction,
)

MINIMAL_TEST_COLUMNS: list[str] = [
    "IndividuCivilite.NomCourt",
    "Individu.CodeAdherent",
    "Individu.Nom",
    "Individu.Prenom",
    "Structure.CodeStructure",
    "Structure.Nom",
    "Fonction.Code",
    "Fonction.CategorieMembre",
    "Fonction.NomMasculin",
    "Fonction.NomFeminin",
    "Inscription.Delegations",
    "Individu.Adresse.Ligne1",
    "Individu.Adresse.Ligne2",
    "Individu.Adresse.Ligne3",
    "Individu.Adresse.CodePostal",
    "Individu.Adresse.Municipalite",
    "Individu.Adresse.Pays",
    "Individu.TelephoneDomicile",
    "Individu.TelephonePortable1",
    "Individu.TelephonePortable2",
    "Individu.CourrielPersonnel",
    "Individu.CourrielDédiéSGDF",
    "RepresentantLegal1Civilite.NomCourt",
    "RepresentantLegal1.CodeAdherent",
    "RepresentantLegal1.Nom",
    "RepresentantLegal1.Prenom",
    "RepresentantLegal1.Adresse.Ligne1",
    "RepresentantLegal1.Adresse.Ligne2",
    "RepresentantLegal1.Adresse.Ligne3",
    "RepresentantLegal1.Adresse.CodePostal",
    "RepresentantLegal1.Adresse.Municipalite",
    "RepresentantLegal1.Adresse.Pays",
    "RepresentantLegal1.TelephoneDomicile",
    "RepresentantLegal1.TelephonePortable1",
    "RepresentantLegal1.TelephonePortable2",
    "RepresentantLegal1.CourrielPersonnel",
    "RepresentantLegal1.CourrielDédiéSGDF",
    "RepresentantLegal2Civilite.NomCourt",
    "RepresentantLegal2.CodeAdherent",
    "RepresentantLegal2.Nom",
    "RepresentantLegal2.Prenom",
    "RepresentantLegal2.Adresse.Ligne1",
    "RepresentantLegal2.Adresse.Ligne2",
    "RepresentantLegal2.Adresse.Ligne3",
    "RepresentantLegal2.Adresse.CodePostal",
    "RepresentantLegal2.Adresse.Municipalite",
    "RepresentantLegal2.Adresse.Pays",
    "RepresentantLegal2.TelephoneDomicile",
    "RepresentantLegal2.TelephonePortable1",
    "RepresentantLegal2.TelephonePortable2",
    "RepresentantLegal2.CourrielPersonnel",
    "RepresentantLegal2.CourrielDédiéSGDF",
    "RepresentantLegal3Civilite.NomCourt",
    "RepresentantLegal3.CodeAdherent",
    "RepresentantLegal3.Nom",
    "RepresentantLegal3.Prenom",
    "RepresentantLegal3.Adresse.Ligne1",
    "RepresentantLegal3.Adresse.Ligne2",
    "RepresentantLegal3.Adresse.Ligne3",
    "RepresentantLegal3.Adresse.CodePostal",
    "RepresentantLegal3.Adresse.Municipalite",
    "RepresentantLegal3.Adresse.Pays",
    "RepresentantLegal3.TelephoneDomicile",
    "RepresentantLegal3.TelephonePortable1",
    "RepresentantLegal3.TelephonePortable2",
    "RepresentantLegal3.CourrielPersonnel",
    "RepresentantLegal3.CourrielDédiéSGDF",
    "Inscriptions.Type",
]


def build_test_row(values: dict[str, str]) -> list[str]:
    """Crée une ligne conforme à `MINIMAL_TEST_COLUMNS` à partir d'un dictionnaire."""
    return [values.get(col, "") for col in MINIMAL_TEST_COLUMNS]


def write_synthetic_html_xls(file_path: Path, rows: list[list[str]]) -> None:
    """Écrit un fichier HTML portant l'extension `.xls` imitant l'export de l'intranet SGDF."""
    header_cells = "".join(f"<th>{col}</th>" for col in MINIMAL_TEST_COLUMNS)
    html_rows = [f"<tr>{header_cells}</tr>"]
    for row in rows:
        cells = "".join(f"<td>{val}</td>" for val in row)
        html_rows.append(f"<tr>{cells}</tr>")
    content = f"<html><body><table>{''.join(html_rows)}</table></body></html>"
    file_path.write_text(content, encoding="utf-8")


def test_normalization_names_phones_units_and_addresses() -> None:
    """Vérifie la normalisation des noms, prénoms composés, téléphones, unités courtes et adresses."""
    assert normalize_last_name("  dupont-moretti ") == "Dupont-Moretti"
    assert normalize_last_name("BERNARD") == "Bernard"
    assert normalize_last_name("  le   goff ") == "Le Goff"
    assert normalize_last_name("d'artagnan") == "D'Artagnan"
    assert normalize_last_name("ÉVRARD") == "Évrard"
    assert normalize_last_name("   ") == ""
    assert normalize_first_name("JEAN-BAPTISTE") == "Jean-Baptiste"
    assert normalize_first_name("  marie   claire ") == "Marie Claire"
    assert normalize_first_name("pierre-yves") == "Pierre-Yves"

    assert normalize_phone("0687966975") == "+33687966975"
    assert normalize_phone("06 87 96 69 75") == "+33687966975"
    assert normalize_phone("05.46.12.34.56") == "+33546123456"
    assert normalize_phone("0033687966975") == "+33687966975"
    assert normalize_phone("+33 (0)6 87 96 69 75") == "+33687966975"
    assert normalize_phone("687966975") == "+33687966975"
    assert normalize_phone("") == ""

    assert normalize_short_unit("1ERE FARFADETS DE LA ROCHELLE", "170") == "Farfadets"
    assert (
        normalize_short_unit("1ERE LOUVETEAUX JEANNETTES DE LA ROCHELLE", "110")
        == "Louveteaux-Jeannettes"
    )
    assert normalize_short_unit("1ERE SCOUTS GUIDES DE LA ROCHELLE", "120") == "Scouts-Guides"
    assert (
        normalize_short_unit("1ERE PIONNIERS CARAVELLES DE LA ROCHELLE", "130")
        == "Pionniers-Caravelles"
    )
    assert normalize_short_unit("3eme COMPAGNONS DE LA ROCHELLE", "140") == "Compagnons"

    street = join_street_lines(
        "10 RUE DU PORT",
        "BÂTIMENT B",
        "",
    )
    assert street == "10 Rue Du Port, Bâtiment B"
    assert normalize_address_part("LA ROCHELLE") == "La Rochelle"
    assert normalize_address_part("SAINT-XANDRE") == "Saint-Xandre"
    assert normalize_address_part("L'HOUMEAU") == "L'Houmeau"
    assert join_street_lines("  ", "", "   ") == ""


def test_sibling_deduplication_enrichment_and_alphabetical_sorting() -> None:
    """Vérifie le dédoublonnage sur (Nom, Prénom), l'enrichissement adulte et le tri accentué."""
    rows = [
        # Jeune 1 : Scouts-Guides (Invité(e)) avec Parent 1 (Élodie BERNARD) et Parent 2 (Marc BERNARD)
        build_test_row(
            {
                "Individu.Nom": "BERNARD",
                "Individu.Prenom": "LUCAS",
                "Structure.Nom": "1ERE SCOUTS GUIDES DE LA ROCHELLE",
                "Fonction.Code": "120",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Invité(e)",
                "RepresentantLegal1.Nom": "bernard",
                "RepresentantLegal1.Prenom": "élodie",
                "RepresentantLegal1.TelephonePortable1": "0601020304",
                "RepresentantLegal1.CourrielPersonnel": "elodie.bernard@example.fr",
                "RepresentantLegal1.Adresse.Ligne1": "5 Avenue Coligny",
                "RepresentantLegal1.Adresse.CodePostal": "17000",
                "RepresentantLegal1.Adresse.Municipalite": "LA ROCHELLE",
                "RepresentantLegal2.Nom": "BERNARD",
                "RepresentantLegal2.Prenom": "MARC",
                "RepresentantLegal2.TelephonePortable1": "0605060708",
            }
        ),
        # Jeune 2 (fratrie) : Farfadets (Inscrit(e)) avec les mêmes parents (ordre inverse de l'âge)
        build_test_row(
            {
                "Individu.Nom": "BERNARD",
                "Individu.Prenom": "CHLOÉ",
                "Structure.Nom": "1ERE FARFADETS DE LA ROCHELLE",
                "Fonction.Code": "170",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Inscrit(e)",
                "RepresentantLegal1.Nom": "BERNARD",
                "RepresentantLegal1.Prenom": "ÉLODIE",
                "RepresentantLegal1.TelephoneDomicile": "0546001122",
                "RepresentantLegal2.Nom": "Bernard",
                "RepresentantLegal2.Prenom": "Marc",
                "RepresentantLegal2.TelephonePortable2": "0700112233",
            }
        ),
        # Jeune 3 : Scouts-Guides avec un parent portant un nom commençant par É (ÉVRARD)
        build_test_row(
            {
                "Individu.Nom": "MARTIN",
                "Individu.Prenom": "TOM",
                "Structure.Nom": "1ERE SCOUTS GUIDES DE LA ROCHELLE",
                "Fonction.Code": "120",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Pré-inscrit(e)",
                "RepresentantLegal1.Nom": "ÉVRARD",
                "RepresentantLegal1.Prenom": "Sophie",
                "RepresentantLegal1.TelephonePortable1": "0611223344",
            }
        ),
        # Jeune 4 : Pionniers-Caravelles avec un parent ZIMMER
        build_test_row(
            {
                "Individu.Nom": "ZIMMER",
                "Individu.Prenom": "HUGO",
                "Structure.Nom": "1ERE PIONNIERS CARAVELLES DE LA ROCHELLE",
                "Fonction.Code": "130",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Inscrit(e)",
                "RepresentantLegal1.Nom": "ZIMMER",
                "RepresentantLegal1.Prenom": "Paul",
                "RepresentantLegal1.TelephonePortable1": "0699887766",
            }
        ),
        # Fiche Adulte pour Élodie BERNARD : Parent Animateur Farfadet avec numéro d'adhérent
        build_test_row(
            {
                "IndividuCivilite.NomCourt": "Mme",
                "Individu.CodeAdherent": "987654321",
                "Individu.Nom": "BERNARD",
                "Individu.Prenom": "Élodie",
                "Structure.Nom": "1ERE FARFADETS DE LA ROCHELLE",
                "Fonction.Code": "271",
                "Fonction.CategorieMembre": "1",
                "Fonction.NomMasculin": "PARENT ANIMATEUR FARFADET",
                "Fonction.NomFeminin": "PARENT ANIMATEUR FARFADET",
            }
        ),
    ]

    contacts = extract_and_group_parents(MINIMAL_TEST_COLUMNS, rows)

    # 4 parents uniques attendus : Bernard Élodie, Bernard Marc, Évrard Sophie, Zimmer Paul
    assert [(c.last_name, c.first_name) for c in contacts] == [
        ("Bernard", "Élodie"),
        ("Bernard", "Marc"),
        ("Évrard", "Sophie"),
        ("Zimmer", "Paul"),
    ]

    elodie = contacts[0]
    assert elodie.children_names == "Lucas Bernard, Chloé Bernard"
    assert elodie.children_units == "Scouts-Guides, Farfadets"
    assert elodie.primary_mobile_phone == "+33601020304"
    assert elodie.landline_phone == "+33546001122"
    assert elodie.street == "5 Avenue Coligny"
    assert elodie.postal_code == "17000"
    assert elodie.city == "La Rochelle"
    assert elodie.member_id == "987654321"
    assert "Parent Animateur Farfadet" in elodie.notes
    assert "Enfant Lucas : Invité(e)" in elodie.notes

    # Pour Sophie Évrard dont l'enfant s'appelle Tom Martin (nom différent), le nom de l'enfant est précisé
    sophie = contacts[2]
    assert sophie.children_names == "Tom Martin"
    assert sophie.children_units == "Scouts-Guides"
    assert sophie.notes == "Enfant Tom : Pré-inscrit(e)"


def test_mobile_phone_promotion_from_landline_or_secondary() -> None:
    """Vérifie le déplacement d'un portable (+336/+337) vers portable 1 avec priorité au fixe."""
    rows = [
        # Cas 1 : Portable 1 vide, Fixe contient un +336 et Portable 2 contient un +337 -> priorité au Fixe
        build_test_row(
            {
                "Individu.Nom": "DURAND",
                "Individu.Prenom": "ALICE",
                "Structure.Nom": "1ERE FARFADETS DE LA ROCHELLE",
                "Fonction.Code": "170",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Inscrit(e)",
                "RepresentantLegal1.Nom": "DURAND",
                "RepresentantLegal1.Prenom": "CLAIRE",
                "RepresentantLegal1.TelephonePortable1": "",
                "RepresentantLegal1.TelephonePortable2": "0711223344",
                "RepresentantLegal1.TelephoneDomicile": "0699887766",
                # Cas 2 : Portable 1 vide, Fixe est un vrai fixe (+335), Portable 2 contient un +337 -> Portable 2 déplacé
                "RepresentantLegal2.Nom": "DURAND",
                "RepresentantLegal2.Prenom": "THOMAS",
                "RepresentantLegal2.TelephonePortable1": "",
                "RepresentantLegal2.TelephonePortable2": "0755667788",
                "RepresentantLegal2.TelephoneDomicile": "0546112233",
            }
        ),
    ]

    contacts = extract_and_group_parents(MINIMAL_TEST_COLUMNS, rows)
    claire, thomas = contacts[0], contacts[1]

    assert claire.first_name == "Claire"
    assert claire.primary_mobile_phone == "+33699887766"
    assert claire.secondary_mobile_phone == "+33711223344"
    assert claire.landline_phone == ""

    assert thomas.first_name == "Thomas"
    assert thomas.primary_mobile_phone == "+33755667788"
    assert thomas.secondary_mobile_phone == ""
    assert thomas.landline_phone == "+33546112233"


def test_excel_workbook_generation_and_cli(tmp_path: Path) -> None:
    """Vérifie la génération complète du fichier `.xlsx` via la CLI."""
    input_path = tmp_path / "ExportTest.xls"
    output_path = tmp_path / "ContactsParentsTest.xlsx"

    write_synthetic_html_xls(
        input_path,
        [
            build_test_row(
                {
                    "Individu.Nom": "DUPONT",
                    "Individu.Prenom": "LÉA",
                    "Structure.Nom": "1ERE FARFADETS DE LA ROCHELLE",
                    "Fonction.Code": "170",
                    "Fonction.CategorieMembre": "0",
                    "Inscriptions.Type": "Inscrit(e)",
                    "RepresentantLegal1.CodeAdherent": "111222333",
                    "RepresentantLegal1.Nom": "DUPONT",
                    "RepresentantLegal1.Prenom": "JEAN-PIERRE",
                    "RepresentantLegal1.TelephonePortable1": "0612345678",
                    "RepresentantLegal1.Adresse.Ligne1": "10 Rue du Port",
                    "RepresentantLegal1.Adresse.Ligne2": "Bâtiment B",
                    "RepresentantLegal1.Adresse.CodePostal": "17000",
                    "RepresentantLegal1.Adresse.Municipalite": "LA ROCHELLE",
                    "RepresentantLegal1.Adresse.Pays": "FRANCE",
                }
            )
        ],
    )

    exit_code = main(["--input", str(input_path), "--output", str(output_path)])
    assert exit_code == 0
    assert output_path.is_file()

    workbook = load_workbook(output_path, data_only=True)
    assert workbook.sheetnames == [SOURCE_SHEET_NAME, PARENTS_SHEET_NAME]

    parents_sheet = workbook[PARENTS_SHEET_NAME]
    headers = [cell.value for cell in parents_sheet[1]]
    assert tuple(headers) == PARENT_CONTACTS_HEADERS
    assert len(headers) == 13

    first_row = [cell.value for cell in parents_sheet[2]]
    assert first_row[0] == "Dupont"
    assert first_row[1] == "Jean-Pierre"
    assert first_row[2] == "+33612345678"  # Format international sans espaces
    assert first_row[6] == "Léa Dupont"
    assert first_row[7] == "Farfadets"
    assert first_row[8] == "10 Rue Du Port, Bâtiment B"
    assert first_row[9] == "17000"
    assert first_row[10] == "La Rochelle"
    assert first_row[11] == "111222333"


def test_real_individuals_export_integration(tmp_path: Path) -> None:
    """Vérifie l'extraction sur le fichier réel `ExportIndividus.xls` sans l'altérer."""
    export_path = PROJECT_ROOT / "ExportIndividus.xls"
    if not export_path.is_file():
        return

    digest_before = hashlib.sha256(export_path.read_bytes()).hexdigest()
    output_path = tmp_path / "ContactsParents.xlsx"

    contacts = run_extraction(input_path=export_path, output_path=output_path)

    digest_after = hashlib.sha256(export_path.read_bytes()).hexdigest()
    assert digest_before == digest_after, "`ExportIndividus.xls` ne doit subir aucune modification."

    # 1. Nombre exact de parents uniques dédoublonnés sur (Nom, Prénom) : 134
    assert len(contacts) == 134
    keys = [(c.last_name, c.first_name) for c in contacts]
    assert len(keys) == len(set(keys)), "Aucun doublon (Nom, Prénom) ne doit subsister."

    # 2. Vérification du formatage Nom et Prénom (casse titre) et téléphones (+33 sans espaces)
    for contact in contacts:
        assert contact.last_name == normalize_last_name(contact.last_name)
        assert contact.last_name != contact.last_name.upper()
        assert contact.first_name == normalize_first_name(contact.first_name)
        assert contact.first_name != contact.first_name.upper()
        assert contact.children_names != ""
        assert contact.children_units != ""
        for phone in (
            contact.primary_mobile_phone,
            contact.secondary_mobile_phone,
            contact.landline_phone,
        ):
            if phone:
                assert phone.startswith("+33")
                assert " " not in phone

    # 3. Vérification de l'enrichissement des Parents Animateurs Farfadets (18 fonctions principales + 2 délégations = 20)
    farfadet_parents = [
        c for c in contacts if "Parent Animateur Farfadet" in c.notes
    ]
    assert len(farfadet_parents) == 20
    assert sum(1 for c in farfadet_parents if c.member_id != "") >= 13

    # 4. Vérification du classeur Excel généré
    workbook = load_workbook(output_path, data_only=True)
    assert workbook.sheetnames == [SOURCE_SHEET_NAME, PARENTS_SHEET_NAME]
    assert workbook[SOURCE_SHEET_NAME].max_row == 154  # 1 en-tête + 153 lignes
    assert workbook[PARENTS_SHEET_NAME].max_row == 135  # 1 en-tête + 134 parents
