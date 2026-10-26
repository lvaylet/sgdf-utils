"""Tests unitaires et d'intégration pour la persistance SQLite (`sgdf_contacts.persistence`)."""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import inspect, text

from sgdf_contacts import (
    PARENTS_SHEET_NAME,
    AttachedChild,
    ParentContact,
    ParentContactRepository,
    create_db_engine,
    extract_and_group_parents,
    format_vcf_content,
    init_db,
    read_individuals_export_table,
    write_parent_contacts_workbook,
)


def _build_sample_parents() -> list[ParentContact]:
    """Crée un jeu synthétique de parents et d'enfants pour tester la persistance."""
    parent_alice = ParentContact(
        last_name="Martin",
        first_name="Alice",
        primary_mobile_phone="+33611223344",
        secondary_mobile_phone="+33699887766",
        landline_phone="+33546112233",
        email_address="alice.martin@example.org",
        street="10 Rue du Port",
        postal_code="17000",
        city="La Rochelle",
        member_id="123456789",
    )
    parent_alice.add_child(
        AttachedChild(
            first_name="Léo",
            last_name="Martin",
            short_unit="LOUVETEAUX-JEANNETTES",
            registration_status="Inscrit(e)",
        )
    )
    parent_alice.add_child(
        AttachedChild(
            first_name="Chloé",
            last_name="Martin",
            short_unit="FARFADETS",
            registration_status="Inscrit(e)",
        )
    )
    parent_alice.add_note("Parent Animateur Farfadet")

    parent_bruno = ParentContact(
        last_name="Évrard",
        first_name="Bruno",
        primary_mobile_phone="",
        secondary_mobile_phone="",
        landline_phone="+33546001122",
        email_address="bruno.evrard@example.org",
        street="5 Avenue des Minimes",
        postal_code="17000",
        city="La Rochelle",
        member_id="",
    )
    parent_bruno.add_child(
        AttachedChild(
            first_name="Emma",
            last_name="Évrard",
            short_unit="SCOUTS-GUIDES",
            registration_status="Pré-inscrit(e)",
        )
    )

    return [parent_alice, parent_bruno]


def test_sqlite_schema_and_pragmas_enabled(tmp_path: Path) -> None:
    """Vérifie que les 3 tables sont créées et que les pragmas WAL et foreign_keys sont actifs."""
    db_file = tmp_path / "test_pragmas.db"
    engine = create_db_engine(db_file)
    init_db(engine)

    inspector = inspect(engine)
    table_names = set(inspector.get_table_names())
    assert {"parents", "children", "parent_child_links"}.issubset(table_names)

    with engine.connect() as connection:
        journal_mode = connection.execute(text("PRAGMA journal_mode;")).scalar()
        foreign_keys = connection.execute(text("PRAGMA foreign_keys;")).scalar()

    assert str(journal_mode).lower() == "wal"
    assert int(foreign_keys) == 1


def test_sync_parents_idempotence_and_diff_report(tmp_path: Path) -> None:
    """Vérifie l'import initial, l'idempotence au ré-import et le calcul détaillé de `ImportDiffReport`."""
    repo = ParentContactRepository(create_db_engine(tmp_path / "contacts.db"))
    sample_parents = _build_sample_parents()

    # 1. Premier import : 2 créations
    report_initial = repo.sync_parents(sample_parents)
    assert len(report_initial.created) == 2
    assert len(report_initial.updated) == 0
    assert len(report_initial.unchanged) == 0
    assert len(report_initial.deactivated) == 0

    # 2. Ré-import identique : idempotent (0 created, 0 updated, 2 unchanged)
    report_repeat = repo.sync_parents(sample_parents)
    assert len(report_repeat.created) == 0
    assert len(report_repeat.updated) == 0
    assert len(report_repeat.unchanged) == 2
    assert len(report_repeat.deactivated) == 0

    # 3. Enrichissement externe de `has_whatsapp` sur Alice et Bruno
    assert repo.set_whatsapp_status("Martin", "Alice", True) is True
    assert repo.set_whatsapp_status("Évrard", "Bruno", False) is True

    # 4. Nouvel import : modification d'Alice (téléphone + unité d'enfant), retrait de Bruno, ajout de Claire
    updated_alice = ParentContact(
        last_name="Martin",
        first_name="Alice",
        primary_mobile_phone="+33655443322",
        secondary_mobile_phone="+33699887766",
        landline_phone="+33546112233",
        email_address="alice.martin@example.org",
        street="10 Rue du Port",
        postal_code="17000",
        city="La Rochelle",
        member_id="123456789",
    )
    updated_alice.add_child(
        AttachedChild(
            first_name="Léo",
            last_name="Martin",
            short_unit="SCOUTS-GUIDES",
            registration_status="Inscrit(e)",
        )
    )
    updated_alice.add_child(
        AttachedChild(
            first_name="Chloé",
            last_name="Martin",
            short_unit="FARFADETS",
            registration_status="Inscrit(e)",
        )
    )
    updated_alice.add_note("Parent Animateur Farfadet")

    new_claire = ParentContact(
        last_name="Dupont",
        first_name="Claire",
        primary_mobile_phone="+33712345678",
        email_address="claire.dupont@example.org",
    )
    new_claire.add_child(
        AttachedChild(
            first_name="Lucas",
            last_name="Dupont",
            short_unit="PIONNIERS-CARAVELLES",
            registration_status="Inscrit(e)",
        )
    )

    report_diff = repo.sync_parents([updated_alice, new_claire])
    assert [p.first_name for p in report_diff.created] == ["Claire"]
    assert len(report_diff.updated) == 1
    assert report_diff.updated[0].parent.first_name == "Alice"
    assert (
        report_diff.updated[0].changes["primary_mobile_phone"].old_value
        == "+33611223344"
    )
    assert (
        report_diff.updated[0].changes["primary_mobile_phone"].new_value
        == "+33655443322"
    )
    assert "children_units" in report_diff.updated[0].changes
    assert len(report_diff.unchanged) == 0
    assert [p.first_name for p in report_diff.deactivated] == ["Bruno"]

    # Vérifie que `has_whatsapp` a été conservé sur Alice (mise à jour) et Bruno (désactivé)
    alice_record = repo.get_parent_record("Martin", "Alice")
    bruno_record = repo.get_parent_record("Évrard", "Bruno")
    assert alice_record is not None and alice_record.has_whatsapp is True
    assert bruno_record is not None
    assert bruno_record.is_active is False
    assert bruno_record.has_whatsapp is False

    # 5. Réactivation de Bruno lors d'un import ultérieur
    report_reactivate = repo.sync_parents([updated_alice, new_claire, sample_parents[1]])
    assert len(report_reactivate.created) == 0
    assert len(report_reactivate.updated) == 1
    assert report_reactivate.updated[0].parent.first_name == "Bruno"
    assert report_reactivate.updated[0].changes["is_active"].old_value is False
    assert report_reactivate.updated[0].changes["is_active"].new_value is True


def test_reconciliation_by_member_id_when_name_changes(tmp_path: Path) -> None:
    """Vérifie que `member_id` permet de réconcilier un parent même si son prénom/nom est corrigé."""
    repo = ParentContactRepository(create_db_engine(tmp_path / "member_id.db"))
    initial_parent = ParentContact(
        last_name="Le Gall",
        first_name="Jean",
        primary_mobile_phone="+33601020304",
        member_id="999888777",
    )
    repo.sync_parents([initial_parent])

    renamed_parent = ParentContact(
        last_name="Le Gall",
        first_name="Jean-Luc",
        primary_mobile_phone="+33601020304",
        member_id="999888777",
    )
    diff = repo.sync_parents([renamed_parent])
    assert len(diff.created) == 0
    assert len(diff.deactivated) == 0
    assert len(diff.updated) == 1
    assert diff.updated[0].changes["first_name"].old_value == "Jean"
    assert diff.updated[0].changes["first_name"].new_value == "Jean-Luc"


def test_query_filtering_and_phone_lookup(tmp_path: Path) -> None:
    """Vérifie le tri alphabétique français, les filtres multicritères et la recherche par téléphone."""
    repo = ParentContactRepository(create_db_engine(tmp_path / "queries.db"))
    repo.sync_parents(_build_sample_parents())

    # Tri alphabétique français sans sensibilité aux accents : Évrard avant Martin
    active_parents = repo.list_active_parents()
    assert [p.last_name for p in active_parents] == ["Évrard", "Martin"]

    # Filtre par unité d'enfant
    farfadets_parents = repo.filter_parents(short_unit="FARFADETS")
    assert [p.last_name for p in farfadets_parents] == ["Martin"]

    # Filtre par statut d'inscription d'enfant
    pre_registered = repo.filter_parents(registration_status="Pré-inscrit(e)")
    assert [p.last_name for p in pre_registered] == ["Évrard"]

    # Filtre par présence d'un numéro de portable valide
    with_mobile = repo.filter_parents(has_valid_mobile=True)
    without_mobile = repo.filter_parents(has_valid_mobile=False)
    assert [p.last_name for p in with_mobile] == ["Martin"]
    assert [p.last_name for p in without_mobile] == ["Évrard"]

    # Recherche par téléphone sur portable 1, portable 2 et fixe (formats normalisé et national)
    found_p1 = repo.find_by_phone("06 11 22 33 44")
    found_p2 = repo.find_by_phone("+33699887766")
    found_landline = repo.find_by_phone("05.46.00.11.22")
    assert found_p1 is not None and found_p1.last_name == "Martin"
    assert found_p2 is not None and found_p2.last_name == "Martin"
    assert found_landline is not None and found_landline.last_name == "Évrard"
    assert repo.find_by_phone("+33600000000") is None


def test_regenerate_vcf_and_excel_from_sqlite_synthetic(tmp_path: Path) -> None:
    """Vérifie que `.vcf` et `.xlsx` peuvent être régénérés directement depuis SQLite."""
    repo = ParentContactRepository(create_db_engine(tmp_path / "exports.db"))
    sample_parents = _build_sample_parents()
    repo.sync_parents(sample_parents)
    sorted_parents = repo.list_active_parents()

    vcf_path = tmp_path / "ContactParents.vcf"
    vcf_from_db = repo.export_vcf(vcf_path)
    vcf_direct = format_vcf_content(sorted_parents)
    assert vcf_from_db == vcf_direct
    assert vcf_path.read_bytes().decode("utf-8") == vcf_direct

    xlsx_path = tmp_path / "ContactsParents.xlsx"
    xlsx_bytes = repo.export_excel(xlsx_path)
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    sheet = wb[PARENTS_SHEET_NAME]
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 3  # 1 en-tête + 2 parents
    assert [cell or "" for cell in rows[1]] == sorted_parents[0].to_excel_row()
    assert [cell or "" for cell in rows[2]] == sorted_parents[1].to_excel_row()


def test_real_export_individus_end_to_end_persistence(tmp_path: Path) -> None:
    """Vérifie de bout en bout l'import de `ExportIndividus.xls`, l'idempotence et la régénération VCF/Excel."""
    export_file = Path(__file__).resolve().parent.parent / "ExportIndividus.xls"
    if not export_file.is_file():
        return

    repo = ParentContactRepository(create_db_engine(tmp_path / "real_sgdf.db"))

    # 1. Import initial depuis ExportIndividus.xls
    first_report = repo.import_individuals_xls(export_file)
    assert len(first_report.created) == 134
    assert len(first_report.updated) == 0
    assert len(first_report.unchanged) == 0
    assert len(first_report.deactivated) == 0

    # 2. Ré-import idempotent du même fichier
    second_report = repo.import_individuals_xls(export_file)
    assert len(second_report.created) == 0
    assert len(second_report.updated) == 0
    assert len(second_report.unchanged) == 134
    assert len(second_report.deactivated) == 0

    # 3. Comparaison stricte des exports VCF et Excel régénérés depuis SQLite vs extraction directe
    headers, rows = read_individuals_export_table(export_file)
    direct_parents = extract_and_group_parents(headers, rows)

    vcf_from_db = repo.export_vcf(tmp_path / "ContactParents.vcf")
    vcf_direct = format_vcf_content(direct_parents)
    assert vcf_from_db == vcf_direct

    xlsx_db_bytes = repo.export_excel(tmp_path / "ContactsParents.xlsx")
    xlsx_direct_bytes = write_parent_contacts_workbook(
        output_path=None,
        parent_contacts=direct_parents,
    )
    wb_db = load_workbook(io.BytesIO(xlsx_db_bytes), data_only=True)
    wb_direct = load_workbook(io.BytesIO(xlsx_direct_bytes), data_only=True)

    rows_db = list(wb_db[PARENTS_SHEET_NAME].iter_rows(values_only=True))
    rows_direct = list(wb_direct[PARENTS_SHEET_NAME].iter_rows(values_only=True))
    assert len(rows_db) == 135  # 1 en-tête + 134 parents
    assert rows_db == rows_direct
