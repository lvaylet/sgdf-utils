"""Tests d'intégration de l'API HTTP REST FastAPI (`sgdf_contacts.entrypoints.api`)."""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from sgdf_contacts import (
    PARENTS_SHEET_NAME,
    ApiSettings,
    AttachedChild,
    ParentContact,
    create_app,
)


def _build_html_xls_bytes(rows: list[dict[str, str]]) -> bytes:
    """Génère un flux HTML au format `ExportIndividus.xls` de l'intranet SGDF."""
    headers = [
        "Structure.Nom",
        "Fonction.CategorieMembre",
        "Fonction.Code",
        "Fonction.NomMasculin",
        "Fonction.NomFeminin",
        "IndividuCivilite.NomCourt",
        "Individu.CodeAdherent",
        "Individu.Nom",
        "Individu.Prenom",
        "Individu.CourrielPersonnel",
        "Individu.CourrielDédiéSGDF",
        "Individu.TelephoneDomicile",
        "Individu.TelephonePortable1",
        "Individu.TelephonePortable2",
        "Individu.Adresse.Ligne1",
        "Individu.Adresse.Ligne2",
        "Individu.Adresse.Ligne3",
        "Individu.Adresse.CodePostal",
        "Individu.Adresse.Municipalite",
        "Inscription.Delegations",
        "Inscriptions.Type",
        "RepresentantLegal1.CodeAdherent",
        "RepresentantLegal1.Nom",
        "RepresentantLegal1.Prenom",
        "RepresentantLegal1.CourrielPersonnel",
        "RepresentantLegal1.TelephoneDomicile",
        "RepresentantLegal1.TelephonePortable1",
        "RepresentantLegal1.TelephonePortable2",
        "RepresentantLegal1.Adresse.Ligne1",
        "RepresentantLegal1.Adresse.Ligne2",
        "RepresentantLegal1.Adresse.Ligne3",
        "RepresentantLegal1.Adresse.CodePostal",
        "RepresentantLegal1.Adresse.Municipalite",
        "RepresentantLegal2.CodeAdherent",
        "RepresentantLegal2.Nom",
        "RepresentantLegal2.Prenom",
        "RepresentantLegal2.CourrielPersonnel",
        "RepresentantLegal2.TelephoneDomicile",
        "RepresentantLegal2.TelephonePortable1",
        "RepresentantLegal2.TelephonePortable2",
        "RepresentantLegal2.Adresse.Ligne1",
        "RepresentantLegal2.Adresse.Ligne2",
        "RepresentantLegal2.Adresse.Ligne3",
        "RepresentantLegal2.Adresse.CodePostal",
        "RepresentantLegal2.Adresse.Municipalite",
        "RepresentantLegal3.Nom",
        "RepresentantLegal3.Prenom",
    ]
    th_cells = "".join(f"<th>{h}</th>" for h in headers)
    tr_rows: list[str] = [f"<tr>{th_cells}</tr>"]
    for row in rows:
        td_cells = "".join(f"<td>{row.get(h, '')}</td>" for h in headers)
        tr_rows.append(f"<tr>{td_cells}</tr>")

    html = f"<html><body><table>{''.join(tr_rows)}</table></body></html>"
    return html.encode("utf-8")


def _sample_parents_for_api() -> list[ParentContact]:
    """Construit un jeu de parents varié pour tester les routes de consultation et WAHA."""
    parent_alice = ParentContact(
        last_name="Martin",
        first_name="Alice",
        primary_mobile_phone="+33687966975",
        secondary_mobile_phone="+33611223344",
        landline_phone="+33546112233",
        email_address="alice.martin@example.org",
        street="10 Rue du Port",
        postal_code="17000",
        city="La Rochelle",
        member_id="123456789",
    )
    parent_alice.add_child(
        AttachedChild(
            first_name="Léo",
            last_name="Martin",
            short_unit="LOUVETEAUX-JEANNETTES",
            registration_status="Inscrit(e)",
        )
    )
    parent_alice.add_child(
        AttachedChild(
            first_name="Chloé",
            last_name="Martin",
            short_unit="FARFADETS",
            registration_status="Inscrit(e)",
        )
    )
    parent_alice.add_note("Parent Animateur Farfadet")

    parent_bruno = ParentContact(
        last_name="Bernard",
        first_name="Bruno",
        primary_mobile_phone="",
        secondary_mobile_phone="+33712345678",
        landline_phone="",
        email_address="bruno.bernard@example.org",
        street="8 Rue des Voiliers",
        postal_code="17000",
        city="La Rochelle",
        member_id="",
    )
    parent_bruno.add_child(
        AttachedChild(
            first_name="Lucas",
            last_name="Bernard",
            short_unit="FARFADETS",
            registration_status="Pré-inscrit(e)",
        )
    )

    parent_claire = ParentContact(
        last_name="Dubois",
        first_name="Claire",
        primary_mobile_phone="",
        secondary_mobile_phone="",
        landline_phone="+33546998877",
        email_address="claire.dubois@example.org",
        street="3 Place de Verdun",
        postal_code="17000",
        city="La Rochelle",
        member_id="",
    )
    parent_claire.add_child(
        AttachedChild(
            first_name="Emma",
            last_name="Dubois",
            short_unit="FARFADETS",
            registration_status="Inscrit(e)",
        )
    )

    return [parent_alice, parent_bruno, parent_claire]


def test_health_and_openapi_endpoints(tmp_path: Path) -> None:
    """Vérifie que `/health`, `/docs` et `/openapi.json` répondent `200` et exposent les routes."""
    app = create_app(database_url=tmp_path / "api_health.db")
    client = TestClient(app)

    health_resp = client.get("/health")
    assert health_resp.status_code == 200
    assert health_resp.json() == {"status": "ok", "database": "ok"}

    openapi_resp = client.get("/openapi.json")
    assert openapi_resp.status_code == 200
    schema = openapi_resp.json()
    paths = schema.get("paths", {})
    assert "/health" in paths
    assert "/api/v1/imports/intranet" in paths
    assert "/api/v1/parents" in paths
    assert "/api/v1/units/{unit_name}/whatsapp-recipients" in paths
    assert "/api/v1/parents/by-phone/{phone}" in paths
    assert "/api/v1/parents/{parent_id}" in paths
    assert "/api/v1/exports/vcf" in paths
    assert "/api/v1/exports/xlsx" in paths

    docs_resp = client.get("/docs")
    assert docs_resp.status_code == 200
    assert "Swagger UI" in docs_resp.text


def test_api_settings_and_x_api_key_authentication(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Vérifie la configuration par variables d'environnement et l'authentification `X-API-Key`."""
    monkeypatch.setenv("SGDF_DATABASE_URL", f"sqlite:///{tmp_path / 'env.db'}")
    monkeypatch.setenv("SGDF_API_KEY", "super-secret-homelab-key")
    monkeypatch.setenv("SGDF_HOST", "0.0.0.0")
    monkeypatch.setenv("SGDF_PORT", "9090")

    settings = ApiSettings.from_env()
    assert settings.api_key == "super-secret-homelab-key"
    assert settings.host == "0.0.0.0"
    assert settings.port == 9090

    app = create_app()
    client = TestClient(app)

    # /health reste accessible sans clé d'API
    assert client.get("/health").status_code == 200

    # Sans en-tête X-API-Key -> 401
    resp_missing = client.get("/api/v1/parents")
    assert resp_missing.status_code == 401
    assert "X-API-Key" in resp_missing.json()["detail"]

    # Avec une mauvaise clé -> 403
    resp_wrong = client.get(
        "/api/v1/parents",
        headers={"X-API-Key": "wrong-key"},
    )
    assert resp_wrong.status_code == 403

    # Avec la bonne clé -> 200
    resp_ok = client.get(
        "/api/v1/parents",
        headers={"X-API-Key": "super-secret-homelab-key"},
    )
    assert resp_ok.status_code == 200
    assert resp_ok.json() == []

    # Lorsque SGDF_API_KEY est supprimée, l'accès sans clé est autorisé
    monkeypatch.delenv("SGDF_API_KEY", raising=False)
    resp_open = client.get("/api/v1/parents")
    assert resp_open.status_code == 200


def test_import_intranet_multipart_and_diff_report(tmp_path: Path) -> None:
    """Vérifie l'upload multipart `POST /api/v1/imports/intranet` et le calcul de différentiel JSON."""
    app = create_app(database_url=tmp_path / "imports.db")
    client = TestClient(app)

    initial_rows = [
        {
            "Structure.Nom": "LOUVETEAUX-JEANNETTES",
            "Fonction.CategorieMembre": "0",
            "Fonction.Code": "110",
            "Individu.CodeAdherent": "990001",
            "Individu.Nom": "MARTIN",
            "Individu.Prenom": "Léo",
            "Inscriptions.Type": "Inscrit(e)",
            "RepresentantLegal1.Nom": "MARTIN",
            "RepresentantLegal1.Prenom": "Alice",
            "RepresentantLegal1.TelephonePortable1": "06 87 96 69 75",
            "RepresentantLegal1.CourrielPersonnel": "alice.martin@example.org",
            "RepresentantLegal1.Adresse.Ligne1": "10 Rue du Port",
            "RepresentantLegal1.Adresse.CodePostal": "17000",
            "RepresentantLegal1.Adresse.Municipalite": "LA ROCHELLE",
            "RepresentantLegal2.Nom": "MARTIN",
            "RepresentantLegal2.Prenom": "Marc",
            "RepresentantLegal2.TelephonePortable1": "06 12 34 56 78",
        }
    ]
    xls_bytes_v1 = _build_html_xls_bytes(initial_rows)

    # 1. Premier import : 2 parents créés
    resp1 = client.post(
        "/api/v1/imports/intranet",
        files={
            "file": (
                "ExportIndividus.xls",
                io.BytesIO(xls_bytes_v1),
                "application/vnd.ms-excel",
            )
        },
    )
    assert resp1.status_code == 200
    data1 = resp1.json()
    assert data1["total_parents"] == 2
    assert data1["created"] == 2
    assert data1["updated"] == 0
    assert data1["unchanged"] == 0
    assert data1["deactivated"] == 0
    assert len(data1["changes"]["created"]) == 2

    # 2. Second import identique : idempotence stricte
    resp2 = client.post(
        "/api/v1/imports/intranet",
        files={
            "file": (
                "ExportIndividus.xls",
                io.BytesIO(xls_bytes_v1),
                "application/vnd.ms-excel",
            )
        },
    )
    assert resp2.status_code == 200
    data2 = resp2.json()
    assert data2["total_parents"] == 2
    assert data2["created"] == 0
    assert data2["updated"] == 0
    assert data2["unchanged"] == 2
    assert data2["deactivated"] == 0

    # 3. Troisième import modifié : Alice change de téléphone, Marc est retiré, Sophie est ajoutée
    updated_rows = [
        {
            "Structure.Nom": "LOUVETEAUX-JEANNETTES",
            "Fonction.CategorieMembre": "0",
            "Fonction.Code": "110",
            "Individu.CodeAdherent": "990001",
            "Individu.Nom": "MARTIN",
            "Individu.Prenom": "Léo",
            "Inscriptions.Type": "Inscrit(e)",
            "RepresentantLegal1.Nom": "MARTIN",
            "RepresentantLegal1.Prenom": "Alice",
            "RepresentantLegal1.TelephonePortable1": "06 99 88 77 66",
            "RepresentantLegal1.CourrielPersonnel": "alice.martin@example.org",
            "RepresentantLegal1.Adresse.Ligne1": "10 Rue du Port",
            "RepresentantLegal1.Adresse.CodePostal": "17000",
            "RepresentantLegal1.Adresse.Municipalite": "LA ROCHELLE",
            "RepresentantLegal2.Nom": "DURAND",
            "RepresentantLegal2.Prenom": "Sophie",
            "RepresentantLegal2.TelephonePortable1": "07 00 11 22 33",
        }
    ]
    xls_bytes_v2 = _build_html_xls_bytes(updated_rows)
    resp3 = client.post(
        "/api/v1/imports/intranet",
        files={
            "file": (
                "ExportIndividus.xls",
                io.BytesIO(xls_bytes_v2),
                "application/vnd.ms-excel",
            )
        },
    )
    assert resp3.status_code == 200
    data3 = resp3.json()
    assert data3["total_parents"] == 2
    assert data3["created"] == 1
    assert data3["updated"] == 1
    assert data3["unchanged"] == 0
    assert data3["deactivated"] == 1

    updated_entry = data3["changes"]["updated"][0]
    assert updated_entry["parent"]["first_name"] == "Alice"
    assert updated_entry["changes"]["primary_mobile_phone"] == {
        "old_value": "+33687966975",
        "new_value": "+33699887766",
    }
    assert data3["changes"]["deactivated"][0]["first_name"] == "Marc"

    # 4. Fichier vide ou invalide -> 400
    resp_empty = client.post(
        "/api/v1/imports/intranet",
        files={"file": ("empty.xls", io.BytesIO(b""), "application/vnd.ms-excel")},
    )
    assert resp_empty.status_code == 400

    resp_invalid = client.post(
        "/api/v1/imports/intranet",
        files={
            "file": (
                "invalid.xls",
                io.BytesIO(b"<html><body><table><tr><th>Col</th></tr></table></body></html>"),
                "application/vnd.ms-excel",
            )
        },
    )
    assert resp_invalid.status_code == 400


def test_whatsapp_recipients_and_by_phone_and_patch(tmp_path: Path) -> None:
    """Vérifie `whatsapp-recipients` (`waha_chat_id`), `by-phone` multi-formats et `PATCH /parents/{id}`."""
    app = create_app(database_url=tmp_path / "waha.db")
    client = TestClient(app)

    # Peuple la base via le dépôt de l'application
    repo = app.state.repository
    assert repo is not None
    repo.sync_parents(_sample_parents_for_api())

    # 1. GET /api/v1/units/{unit_name}/whatsapp-recipients
    # Sur FARFADETS : Alice (mobile principal +33687966975) et Bruno (mobile secondaire +33712345678)
    # Claire n'a qu'un fixe (+33546998877) donc elle est exclue des destinataires WhatsApp
    recipients_resp = client.get("/api/v1/units/FARFADETS/whatsapp-recipients")
    assert recipients_resp.status_code == 200
    recipients = recipients_resp.json()
    assert len(recipients) == 2

    # Tri alphabétique : Bernard Bruno puis Martin Alice
    bruno_recipient = recipients[0]
    assert bruno_recipient["last_name"] == "Bernard"
    assert bruno_recipient["first_name"] == "Bruno"
    assert bruno_recipient["phone_e164"] == "+33712345678"
    assert bruno_recipient["waha_chat_id"] == "33712345678@c.us"
    assert len(bruno_recipient["children"]) == 1

    alice_recipient = recipients[1]
    assert alice_recipient["last_name"] == "Martin"
    assert alice_recipient["first_name"] == "Alice"
    assert alice_recipient["phone_e164"] == "+33687966975"
    assert alice_recipient["waha_chat_id"] == "33687966975@c.us"
    assert len(alice_recipient["children"]) == 2

    # Recherche par code court d'unité (LJ -> LOUVETEAUX-JEANNETTES)
    lj_resp = client.get("/api/v1/units/LJ/whatsapp-recipients")
    assert lj_resp.status_code == 200
    lj_recipients = lj_resp.json()
    assert len(lj_recipients) == 1
    assert lj_recipients[0]["waha_chat_id"] == "33687966975@c.us"

    # 2. GET /api/v1/parents/by-phone/{phone} sous tous les formats requis
    for phone_variant in (
        "+33687966975",
        "0687966975",
        "33687966975",
        "33687966975@c.us",
    ):
        found_resp = client.get(f"/api/v1/parents/by-phone/{phone_variant}")
        assert found_resp.status_code == 200, f"Échec pour le format {phone_variant}"
        found_parent = found_resp.json()
        assert found_parent["first_name"] == "Alice"
        assert found_parent["last_name"] == "Martin"
        assert found_parent["primary_mobile_phone"] == "+33687966975"

    # Recherche sur le mobile secondaire ou le fixe
    found_secondary = client.get("/api/v1/parents/by-phone/33712345678@c.us")
    assert found_secondary.status_code == 200
    assert found_secondary.json()["first_name"] == "Bruno"

    found_landline = client.get("/api/v1/parents/by-phone/0546998877")
    assert found_landline.status_code == 200
    assert found_landline.json()["first_name"] == "Claire"

    # Numéro inconnu -> 404
    not_found = client.get("/api/v1/parents/by-phone/33600000000@c.us")
    assert not_found.status_code == 404

    # 3. PATCH /api/v1/parents/{parent_id}
    alice_id = alice_recipient["parent_id"]
    patch_resp = client.patch(
        f"/api/v1/parents/{alice_id}",
        json={"has_whatsapp": True},
    )
    assert patch_resp.status_code == 200
    assert patch_resp.json()["has_whatsapp"] is True

    # Vérifie que la valeur mise à jour est visible dans whatsapp-recipients
    lj_after_patch = client.get("/api/v1/units/LJ/whatsapp-recipients").json()
    assert lj_after_patch[0]["has_whatsapp"] is True

    # Identifiant inexistant -> 404
    patch_missing = client.patch(
        "/api/v1/parents/99999",
        json={"has_whatsapp": False},
    )
    assert patch_missing.status_code == 404


def test_list_parents_filters_pagination_and_exports(tmp_path: Path) -> None:
    """Vérifie `GET /api/v1/parents` (filtres, pagination) et les exports `/exports/vcf` et `/exports/xlsx`."""
    app = create_app(database_url=tmp_path / "exports.db")
    client = TestClient(app)

    repo = app.state.repository
    assert repo is not None
    repo.sync_parents(_sample_parents_for_api())

    # Tous les parents actifs (3)
    all_resp = client.get("/api/v1/parents")
    assert all_resp.status_code == 200
    all_parents = all_resp.json()
    assert [p["last_name"] for p in all_parents] == ["Bernard", "Dubois", "Martin"]

    # Pagination : limit=1, offset=1 -> Dubois Claire
    paged_resp = client.get("/api/v1/parents?limit=1&offset=1")
    assert paged_resp.status_code == 200
    assert len(paged_resp.json()) == 1
    assert paged_resp.json()[0]["last_name"] == "Dubois"

    # Filtre has_mobile=true -> Bernard et Martin (2)
    mobile_resp = client.get("/api/v1/parents?has_mobile=true")
    assert mobile_resp.status_code == 200
    assert [p["last_name"] for p in mobile_resp.json()] == ["Bernard", "Martin"]

    # Filtre registration_status=Pré-inscrit(e) -> Bernard Bruno (1)
    status_resp = client.get("/api/v1/parents?registration_status=Pr%C3%A9-inscrit(e)")
    assert status_resp.status_code == 200
    assert [p["last_name"] for p in status_resp.json()] == ["Bernard"]

    # Export VCF à la volée
    vcf_resp = client.get("/api/v1/exports/vcf")
    assert vcf_resp.status_code == 200
    assert "text/vcard" in vcf_resp.headers["content-type"]
    assert "ContactParents.vcf" in vcf_resp.headers["content-disposition"]
    assert vcf_resp.text.count("BEGIN:VCARD") == 3

    # Export VCF filtré par unité LJ (1 parent : Alice Martin)
    vcf_lj_resp = client.get("/api/v1/exports/vcf?unit=LJ")
    assert vcf_lj_resp.status_code == 200
    assert vcf_lj_resp.text.count("BEGIN:VCARD") == 1
    assert "FN:Alice Martin" in vcf_lj_resp.text

    # Export XLSX à la volée
    xlsx_resp = client.get("/api/v1/exports/xlsx")
    assert xlsx_resp.status_code == 200
    assert (
        xlsx_resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "ContactsParents.xlsx" in xlsx_resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(xlsx_resp.content))
    ws = wb[PARENTS_SHEET_NAME]
    assert ws.max_row == 4  # 1 ligne d'en-tête + 3 parents


def test_real_export_individus_end_to_end_via_api(tmp_path: Path) -> None:
    """Vérifie le cycle complet sur le fichier réel `ExportIndividus.xls` (134 parents) s'il est présent."""
    export_path = PROJECT_ROOT / "ExportIndividus.xls"
    if not export_path.is_file():
        return

    app = create_app(
        database_url=tmp_path / "real_intranet.db",
        api_key="n8n-api-key",
    )
    client = TestClient(app)
    headers = {"X-API-Key": "n8n-api-key"}

    raw_bytes = export_path.read_bytes()
    import_resp = client.post(
        "/api/v1/imports/intranet",
        headers=headers,
        files={
            "file": (
                "ExportIndividus.xls",
                io.BytesIO(raw_bytes),
                "application/vnd.ms-excel",
            )
        },
    )
    assert import_resp.status_code == 200
    report = import_resp.json()
    assert report["total_parents"] == 134
    assert report["created"] == 134
    assert report["updated"] == 0
    assert report["unchanged"] == 0
    assert report["deactivated"] == 0

    # Second import idempotent
    import_resp_2 = client.post(
        "/api/v1/imports/intranet",
        headers=headers,
        files={
            "file": (
                "ExportIndividus.xls",
                io.BytesIO(raw_bytes),
                "application/vnd.ms-excel",
            )
        },
    )
    assert import_resp_2.status_code == 200
    report_2 = import_resp_2.json()
    assert report_2["total_parents"] == 134
    assert report_2["created"] == 0
    assert report_2["unchanged"] == 134

    # Vérifie que les exports VCF et XLSX contiennent les 134 parents
    vcf_resp = client.get("/api/v1/exports/vcf", headers=headers)
    assert vcf_resp.status_code == 200
    assert vcf_resp.text.count("BEGIN:VCARD") == 134

    xlsx_resp = client.get("/api/v1/exports/xlsx", headers=headers)
    assert xlsx_resp.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx_resp.content))
    ws = wb[PARENTS_SHEET_NAME]
    assert ws.max_row == 135  # 1 en-tête + 134 parents
