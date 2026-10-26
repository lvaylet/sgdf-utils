"""Tests unitaires et d'intégration pour le package modulaire `sgdf_contacts`.

Vérifie :
1. L'absence de duplication de constantes, d'entités et de logique entre
   `extract_parent_contacts.py`, `generate_contacts_vcf.py` et `sgdf_contacts`.
2. La propriété `waha_chat_id` sur `ParentContact` (`+33687966975` -> `33687966975@c.us`).
3. La lecture de l'export HTML/XLS (`read_individuals_export_table`) depuis des `bytes`
   et un `io.BytesIO` en mémoire en plus d'un `Path`.
4. L'écriture du classeur Excel (`write_parent_contacts_workbook`) dans un flux
   `io.BytesIO` en mémoire.
5. La génération directe de contenu et de fichier `.vcf` à partir d'une `list[ParentContact]`
   en mémoire sans passer par un fichier `.xlsx` intermédiaire.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import extract_parent_contacts as cli_extract
import generate_contacts_vcf as cli_vcf
from sgdf_contacts.domain import (
    ADULT_ROLE_LABELS,
    DEFAULT_CATEGORIES,
    PARENT_CONTACTS_HEADERS,
    PARENTS_SHEET_NAME,
    SOURCE_SHEET_NAME,
    AdultGroupProfile,
    AttachedChild,
    ContactExtractionError,
    ParentContact,
    ParentContactCard,
    VcfGenerationError,
    normalize_first_name,
    normalize_last_name,
    normalize_phone,
)
from sgdf_contacts.exporters import (
    format_parent_vcard,
    format_vcf_content,
    write_parent_contacts_workbook,
    write_vcf_file,
)
from sgdf_contacts.parsers import (
    extract_and_group_parents,
    read_individuals_export_table,
    read_parent_contacts_excel,
)
from tests.test_extract_parent_contacts import MINIMAL_TEST_COLUMNS, build_test_row


def build_synthetic_html_bytes(rows: list[list[str]]) -> bytes:
    """Construit le flux `bytes` UTF-8 d'un export HTML/XLS SGDF synthétique."""
    header_cells = "".join(f"<th>{col}</th>" for col in MINIMAL_TEST_COLUMNS)
    html_rows = [f"<tr>{header_cells}</tr>"]
    for row in rows:
        cells = "".join(f"<td>{val}</td>" for val in row)
        html_rows.append(f"<tr>{cells}</tr>")
    content = f"<html><body><table>{''.join(html_rows)}</table></body></html>"
    return content.encode("utf-8")


def test_no_duplication_between_cli_scripts_and_package() -> None:
    """Vérifie qu'aucune constante, entité ou fonction métier n'est dupliquée entre les scripts CLI."""
    assert cli_extract.PARENT_CONTACTS_HEADERS is PARENT_CONTACTS_HEADERS
    assert cli_vcf.PARENT_CONTACTS_HEADERS is PARENT_CONTACTS_HEADERS
    assert cli_extract.PARENTS_SHEET_NAME is PARENTS_SHEET_NAME
    assert cli_vcf.PARENTS_SHEET_NAME is PARENTS_SHEET_NAME
    assert cli_extract.SOURCE_SHEET_NAME is SOURCE_SHEET_NAME
    assert cli_extract.ADULT_ROLE_LABELS is ADULT_ROLE_LABELS
    assert cli_vcf.DEFAULT_CATEGORIES is DEFAULT_CATEGORIES

    assert cli_extract.ParentContact is ParentContact
    assert cli_extract.AttachedChild is AttachedChild
    assert cli_extract.AdultGroupProfile is AdultGroupProfile
    assert cli_vcf.ParentContactCard is ParentContactCard
    assert cli_extract.ContactExtractionError is ContactExtractionError
    assert cli_vcf.VcfGenerationError is VcfGenerationError

    assert cli_extract.normalize_last_name is normalize_last_name
    assert cli_extract.normalize_first_name is normalize_first_name
    assert cli_extract.normalize_phone is normalize_phone
    assert cli_extract.read_individuals_export_table is read_individuals_export_table
    assert cli_extract.extract_and_group_parents is extract_and_group_parents
    assert cli_extract.write_parent_contacts_workbook is write_parent_contacts_workbook
    assert cli_vcf.read_parent_contacts_excel is read_parent_contacts_excel
    assert cli_vcf.format_parent_vcard is format_parent_vcard
    assert cli_vcf.write_vcf_file is write_vcf_file


def test_parent_contact_waha_chat_id_property() -> None:
    """Vérifie la conversion du numéro portable principal en identifiant WAHA `33687966975@c.us`."""
    contact_with_mobile = ParentContact(
        last_name="Bernard",
        first_name="Élodie",
        primary_mobile_phone="+33687966975",
    )
    assert contact_with_mobile.waha_chat_id == "33687966975@c.us"

    contact_without_plus = ParentContact(
        last_name="Martin",
        first_name="Paul",
        primary_mobile_phone="33612345678",
    )
    assert contact_without_plus.waha_chat_id == "33612345678@c.us"

    contact_without_mobile = ParentContact(
        last_name="Durand",
        first_name="Marie",
        landline_phone="+33546123456",
    )
    assert contact_without_mobile.waha_chat_id == ""


def test_read_individuals_export_table_from_bytes_and_bytesio() -> None:
    """Vérifie que le parseur HTML/XLS lit directement des `bytes` et un `io.BytesIO` en mémoire."""
    raw_rows = [
        build_test_row(
            {
                "Individu.Nom": "BERNARD",
                "Individu.Prenom": "LUCAS",
                "Structure.Nom": "1ERE SCOUTS GUIDES DE LA ROCHELLE",
                "Fonction.Code": "120",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Inscrit(e)",
                "RepresentantLegal1.Nom": "BERNARD",
                "RepresentantLegal1.Prenom": "ÉLODIE",
                "RepresentantLegal1.TelephonePortable1": "0687966975",
                "RepresentantLegal1.CourrielPersonnel": "elodie.bernard@example.fr",
            }
        )
    ]
    html_bytes = build_synthetic_html_bytes(raw_rows)

    headers_from_bytes, rows_from_bytes = read_individuals_export_table(html_bytes)
    headers_from_stream, rows_from_stream = read_individuals_export_table(
        io.BytesIO(html_bytes)
    )

    assert headers_from_bytes == MINIMAL_TEST_COLUMNS
    assert headers_from_stream == MINIMAL_TEST_COLUMNS
    assert len(rows_from_bytes) == 1
    assert rows_from_bytes == rows_from_stream

    parents = extract_and_group_parents(headers_from_bytes, rows_from_bytes)
    assert len(parents) == 1
    assert parents[0].last_name == "Bernard"
    assert parents[0].first_name == "Élodie"
    assert parents[0].waha_chat_id == "33687966975@c.us"

    with pytest.raises(ContactExtractionError, match="Aucun tableau HTML"):
        read_individuals_export_table(b"<html><body>Vide</body></html>")

    with pytest.raises(ContactExtractionError, match="Individu.Nom"):
        read_individuals_export_table(
            b"<html><body><table><tr><th>AutreColonne</th></tr></table></body></html>"
        )


def test_write_workbook_in_memory_and_direct_vcf_from_parent_contacts(
    tmp_path: Path,
) -> None:
    """Vérifie l'écriture Excel en mémoire (`BytesIO`) et la génération vCard directe depuis `list[ParentContact]`."""
    raw_rows = [
        build_test_row(
            {
                "Individu.Nom": "BERNARD",
                "Individu.Prenom": "LUCAS",
                "Structure.Nom": "1ERE SCOUTS GUIDES DE LA ROCHELLE",
                "Fonction.Code": "120",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Invité(e)",
                "RepresentantLegal1.Nom": "BERNARD",
                "RepresentantLegal1.Prenom": "ÉLODIE",
                "RepresentantLegal1.TelephonePortable1": "0687966975",
                "RepresentantLegal1.TelephoneDomicile": "0546001122",
                "RepresentantLegal1.CourrielPersonnel": "elodie.bernard@example.fr",
                "RepresentantLegal1.Adresse.Ligne1": "5 Avenue Coligny",
                "RepresentantLegal1.Adresse.CodePostal": "17000",
                "RepresentantLegal1.Adresse.Municipalite": "LA ROCHELLE",
            }
        ),
        build_test_row(
            {
                "IndividuCivilite.NomCourt": "Mme",
                "Individu.CodeAdherent": "987654321",
                "Individu.Nom": "BERNARD",
                "Individu.Prenom": "Élodie",
                "Structure.Nom": "1ERE FARFADETS DE LA ROCHELLE",
                "Fonction.Code": "271",
                "Fonction.CategorieMembre": "1",
            }
        ),
    ]
    html_bytes = build_synthetic_html_bytes(raw_rows)
    headers, rows = read_individuals_export_table(html_bytes)
    parents = extract_and_group_parents(headers, rows)

    # 1. Écriture du classeur Excel dans un io.BytesIO en mémoire
    memory_stream = io.BytesIO()
    returned_bytes = write_parent_contacts_workbook(
        memory_stream,
        source_headers=headers,
        source_rows=rows,
        parent_contacts=parents,
    )
    assert len(returned_bytes) > 0
    assert memory_stream.getvalue() == returned_bytes

    workbook = load_workbook(memory_stream, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == [SOURCE_SHEET_NAME, PARENTS_SHEET_NAME]
        sheet_rows = list(workbook[PARENTS_SHEET_NAME].iter_rows(values_only=True))
        assert tuple(sheet_rows[0]) == PARENT_CONTACTS_HEADERS
        assert sheet_rows[1][0] == "Bernard"
        assert sheet_rows[1][1] == "Élodie"
    finally:
        workbook.close()

    # 2. Génération vCard directe depuis list[ParentContact] vs passage par un fichier .xlsx
    xlsx_path = tmp_path / "ContactsParents.xlsx"
    vcf_from_excel_path = tmp_path / "FromExcel.vcf"
    vcf_direct_path = tmp_path / "DirectFromParents.vcf"

    xlsx_path.write_bytes(returned_bytes)
    cards_from_excel = read_parent_contacts_excel(xlsx_path)
    write_vcf_file(vcf_from_excel_path, cards_from_excel)
    write_vcf_file(vcf_direct_path, parents)

    direct_vcf_bytes = vcf_direct_path.read_bytes()
    excel_vcf_bytes = vcf_from_excel_path.read_bytes()

    assert direct_vcf_bytes == excel_vcf_bytes
    assert format_vcf_content(parents).encode("utf-8") == excel_vcf_bytes
    direct_vcf_text = direct_vcf_bytes.decode("utf-8")
    assert "FN:Élodie Bernard" in direct_vcf_text
    assert "TEL;TYPE=CELL,VOICE,PREF:+33687966975" in direct_vcf_text
    assert "CATEGORIES:SGDF,La Rochelle" in direct_vcf_text
