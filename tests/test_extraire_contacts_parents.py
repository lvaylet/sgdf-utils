"""Tests unitaires et d'intégration pour `extraire_contacts_parents.py`.

Vérifie :
1. La normalisation des noms (MAJUSCULES), prénoms (casse titre et composés) et unités courtes.
2. Le dédoublonnage des parents sur `(Nom, Prénom)` et le regroupement des fratries et unités.
3. L'enrichissement du numéro d'adhérent et des remarques (rôles adultes comme Parent Animateur
   Farfadet, délégations et statuts d'enfants `Pré-inscrit(e)` / `Invité(e)`).
4. Le tri alphabétique français (gestion des caractères accentués) sur `(Nom, Prénom)`.
5. La création du classeur `ContactsParents.xlsx` avec les onglets `ExportIndividus` et
   `ContactsParents`, ainsi que la lecture seule stricte de `ExportIndividus.xls`.
"""

from __future__ import annotations

import hashlib
import sys
from pathlib import Path

from openpyxl import load_workbook

RACINE_PROJET = Path(__file__).resolve().parent.parent
if str(RACINE_PROJET) not in sys.path:
    sys.path.insert(0, str(RACINE_PROJET))

from extraire_contacts_parents import (
    EN_TETES_CONTACTS_PARENTS,
    NOM_ONGLET_ORIGINE,
    NOM_ONGLET_PARENTS,
    assembler_adresse_physique,
    executer_extraction,
    extraire_et_regrouper_parents,
    normaliser_nom,
    normaliser_prenom,
    normaliser_telephone,
    normaliser_unite_courte,
    principale,
)

COLONNES_MINIMALES_TEST: list[str] = [
    "IndividuCivilite.NomCourt",
    "Individu.CodeAdherent",
    "Individu.Nom",
    "Individu.Prenom",
    "Structure.CodeStructure",
    "Structure.Nom",
    "Fonction.Code",
    "Fonction.CategorieMembre",
    "Fonction.NomMasculin",
    "Fonction.NomFeminin",
    "Inscription.Delegations",
    "Individu.Adresse.Ligne1",
    "Individu.Adresse.Ligne2",
    "Individu.Adresse.Ligne3",
    "Individu.Adresse.CodePostal",
    "Individu.Adresse.Municipalite",
    "Individu.Adresse.Pays",
    "Individu.TelephoneDomicile",
    "Individu.TelephonePortable1",
    "Individu.TelephonePortable2",
    "Individu.CourrielPersonnel",
    "Individu.CourrielDédiéSGDF",
    "RepresentantLegal1Civilite.NomCourt",
    "RepresentantLegal1.CodeAdherent",
    "RepresentantLegal1.Nom",
    "RepresentantLegal1.Prenom",
    "RepresentantLegal1.Adresse.Ligne1",
    "RepresentantLegal1.Adresse.Ligne2",
    "RepresentantLegal1.Adresse.Ligne3",
    "RepresentantLegal1.Adresse.CodePostal",
    "RepresentantLegal1.Adresse.Municipalite",
    "RepresentantLegal1.Adresse.Pays",
    "RepresentantLegal1.TelephoneDomicile",
    "RepresentantLegal1.TelephonePortable1",
    "RepresentantLegal1.TelephonePortable2",
    "RepresentantLegal1.CourrielPersonnel",
    "RepresentantLegal1.CourrielDédiéSGDF",
    "RepresentantLegal2Civilite.NomCourt",
    "RepresentantLegal2.CodeAdherent",
    "RepresentantLegal2.Nom",
    "RepresentantLegal2.Prenom",
    "RepresentantLegal2.Adresse.Ligne1",
    "RepresentantLegal2.Adresse.Ligne2",
    "RepresentantLegal2.Adresse.Ligne3",
    "RepresentantLegal2.Adresse.CodePostal",
    "RepresentantLegal2.Adresse.Municipalite",
    "RepresentantLegal2.Adresse.Pays",
    "RepresentantLegal2.TelephoneDomicile",
    "RepresentantLegal2.TelephonePortable1",
    "RepresentantLegal2.TelephonePortable2",
    "RepresentantLegal2.CourrielPersonnel",
    "RepresentantLegal2.CourrielDédiéSGDF",
    "RepresentantLegal3Civilite.NomCourt",
    "RepresentantLegal3.CodeAdherent",
    "RepresentantLegal3.Nom",
    "RepresentantLegal3.Prenom",
    "RepresentantLegal3.Adresse.Ligne1",
    "RepresentantLegal3.Adresse.Ligne2",
    "RepresentantLegal3.Adresse.Ligne3",
    "RepresentantLegal3.Adresse.CodePostal",
    "RepresentantLegal3.Adresse.Municipalite",
    "RepresentantLegal3.Adresse.Pays",
    "RepresentantLegal3.TelephoneDomicile",
    "RepresentantLegal3.TelephonePortable1",
    "RepresentantLegal3.TelephonePortable2",
    "RepresentantLegal3.CourrielPersonnel",
    "RepresentantLegal3.CourrielDédiéSGDF",
    "Inscriptions.Type",
]


def creer_ligne_test(valeurs: dict[str, str]) -> list[str]:
    """Crée une ligne conforme à `COLONNES_MINIMALES_TEST` à partir d'un dictionnaire."""
    return [valeurs.get(col, "") for col in COLONNES_MINIMALES_TEST]


def ecrire_html_xls_synthetique(chemin: Path, lignes: list[list[str]]) -> None:
    """Écrit un fichier HTML portant l'extension `.xls` imitant l'export de l'intranet SGDF."""
    cellules_entete = "".join(f"<th>{col}</th>" for col in COLONNES_MINIMALES_TEST)
    lignes_html = [f"<tr>{cellules_entete}</tr>"]
    for ligne in lignes:
        cellules = "".join(f"<td>{val}</td>" for val in ligne)
        lignes_html.append(f"<tr>{cellules}</tr>")
    contenu = f"<html><body><table>{''.join(lignes_html)}</table></body></html>"
    chemin.write_text(contenu, encoding="utf-8")


def test_normalisation_noms_prenoms_unites_et_adresses() -> None:
    """Vérifie la normalisation des noms, prénoms composés, téléphones, unités courtes et adresses."""
    assert normaliser_nom("  dupont-moretti ") == "DUPONT-MORETTI"
    assert normaliser_prenom("JEAN-BAPTISTE") == "Jean-Baptiste"
    assert normaliser_prenom("  marie   claire ") == "Marie Claire"
    assert normaliser_prenom("pierre-yves") == "Pierre-Yves"

    assert normaliser_telephone("0687966975") == "+33687966975"
    assert normaliser_telephone("06 87 96 69 75") == "+33687966975"
    assert normaliser_telephone("05.46.12.34.56") == "+33546123456"
    assert normaliser_telephone("0033687966975") == "+33687966975"
    assert normaliser_telephone("+33 (0)6 87 96 69 75") == "+33687966975"
    assert normaliser_telephone("687966975") == "+33687966975"
    assert normaliser_telephone("") == ""

    assert normaliser_unite_courte("1ERE FARFADETS DE LA ROCHELLE", "170") == "Farfadets"
    assert (
        normaliser_unite_courte("1ERE LOUVETEAUX JEANNETTES DE LA ROCHELLE", "110")
        == "Louveteaux-Jeannettes"
    )
    assert normaliser_unite_courte("1ERE SCOUTS GUIDES DE LA ROCHELLE", "120") == "Scouts-Guides"
    assert (
        normaliser_unite_courte("1ERE PIONNIERS CARAVELLES DE LA ROCHELLE", "130")
        == "Pionniers-Caravelles"
    )
    assert normaliser_unite_courte("3eme COMPAGNONS DE LA ROCHELLE", "140") == "Compagnons"

    adresse = assembler_adresse_physique(
        "10 Rue du Port",
        "Bâtiment B",
        "",
        "17000",
        "La Rochelle",
        "FRANCE",
    )
    assert adresse == "10 Rue du Port, Bâtiment B, 17000 LA ROCHELLE"


def test_dedoublonnage_fratrie_enrichissement_et_tri_alphabetique() -> None:
    """Vérifie le dédoublonnage sur (Nom, Prénom), l'enrichissement adulte et le tri accentué."""
    lignes = [
        # Jeune 1 : Scouts-Guides (Invité(e)) avec Parent 1 (Élodie BERNARD) et Parent 2 (Marc BERNARD)
        creer_ligne_test(
            {
                "Individu.Nom": "BERNARD",
                "Individu.Prenom": "LUCAS",
                "Structure.Nom": "1ERE SCOUTS GUIDES DE LA ROCHELLE",
                "Fonction.Code": "120",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Invité(e)",
                "RepresentantLegal1.Nom": "bernard",
                "RepresentantLegal1.Prenom": "élodie",
                "RepresentantLegal1.TelephonePortable1": "0601020304",
                "RepresentantLegal1.CourrielPersonnel": "elodie.bernard@example.fr",
                "RepresentantLegal1.Adresse.Ligne1": "5 Avenue Coligny",
                "RepresentantLegal1.Adresse.CodePostal": "17000",
                "RepresentantLegal1.Adresse.Municipalite": "LA ROCHELLE",
                "RepresentantLegal2.Nom": "BERNARD",
                "RepresentantLegal2.Prenom": "MARC",
                "RepresentantLegal2.TelephonePortable1": "0605060708",
            }
        ),
        # Jeune 2 (fratrie) : Farfadets (Inscrit(e)) avec les mêmes parents (ordre inverse de l'âge)
        creer_ligne_test(
            {
                "Individu.Nom": "BERNARD",
                "Individu.Prenom": "CHLOÉ",
                "Structure.Nom": "1ERE FARFADETS DE LA ROCHELLE",
                "Fonction.Code": "170",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Inscrit(e)",
                "RepresentantLegal1.Nom": "BERNARD",
                "RepresentantLegal1.Prenom": "ÉLODIE",
                "RepresentantLegal1.TelephoneDomicile": "0546001122",
                "RepresentantLegal2.Nom": "Bernard",
                "RepresentantLegal2.Prenom": "Marc",
                "RepresentantLegal2.TelephonePortable2": "0700112233",
            }
        ),
        # Jeune 3 : Scouts-Guides avec un parent portant un nom commençant par É (ÉVRARD)
        creer_ligne_test(
            {
                "Individu.Nom": "MARTIN",
                "Individu.Prenom": "TOM",
                "Structure.Nom": "1ERE SCOUTS GUIDES DE LA ROCHELLE",
                "Fonction.Code": "120",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Pré-inscrit(e)",
                "RepresentantLegal1.Nom": "ÉVRARD",
                "RepresentantLegal1.Prenom": "Sophie",
                "RepresentantLegal1.TelephonePortable1": "0611223344",
            }
        ),
        # Jeune 4 : Pionniers-Caravelles avec un parent ZIMMER
        creer_ligne_test(
            {
                "Individu.Nom": "ZIMMER",
                "Individu.Prenom": "HUGO",
                "Structure.Nom": "1ERE PIONNIERS CARAVELLES DE LA ROCHELLE",
                "Fonction.Code": "130",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Inscrit(e)",
                "RepresentantLegal1.Nom": "ZIMMER",
                "RepresentantLegal1.Prenom": "Paul",
                "RepresentantLegal1.TelephonePortable1": "0699887766",
            }
        ),
        # Fiche Adulte pour Élodie BERNARD : Parent Animateur Farfadet avec numéro d'adhérent
        creer_ligne_test(
            {
                "IndividuCivilite.NomCourt": "Mme",
                "Individu.CodeAdherent": "987654321",
                "Individu.Nom": "BERNARD",
                "Individu.Prenom": "Élodie",
                "Structure.Nom": "1ERE FARFADETS DE LA ROCHELLE",
                "Fonction.Code": "271",
                "Fonction.CategorieMembre": "1",
                "Fonction.NomMasculin": "PARENT ANIMATEUR FARFADET",
                "Fonction.NomFeminin": "PARENT ANIMATEUR FARFADET",
            }
        ),
    ]

    contacts = extraire_et_regrouper_parents(COLONNES_MINIMALES_TEST, lignes)

    # 4 parents uniques attendus : BERNARD Élodie, BERNARD Marc, ÉVRARD Sophie, ZIMMER Paul
    assert [(c.nom, c.prenom) for c in contacts] == [
        ("BERNARD", "Élodie"),
        ("BERNARD", "Marc"),
        ("ÉVRARD", "Sophie"),
        ("ZIMMER", "Paul"),
    ]

    elodie = contacts[0]
    assert elodie.noms_enfants == "Lucas BERNARD, Chloé BERNARD"
    assert elodie.unites_enfants == "Scouts-Guides, Farfadets"
    assert elodie.telephone_portable_principal == "+33601020304"
    assert elodie.telephone_fixe == "+33546001122"
    assert elodie.numero_adherent == "987654321"
    assert "Parent Animateur Farfadet" in elodie.remarques
    assert "Enfant Lucas : Invité(e)" in elodie.remarques

    # Pour Sophie ÉVRARD dont l'enfant s'appelle Tom MARTIN (nom différent), le nom de l'enfant est précisé
    sophie = contacts[2]
    assert sophie.noms_enfants == "Tom MARTIN"
    assert sophie.unites_enfants == "Scouts-Guides"
    assert sophie.remarques == "Enfant Tom : Pré-inscrit(e)"


def test_reorganisation_numero_portable_depuis_fixe_ou_portable2() -> None:
    """Vérifie le déplacement d'un portable (+336/+337) vers portable 1 avec priorité au fixe."""
    lignes = [
        # Cas 1 : Portable 1 vide, Fixe contient un +336 et Portable 2 contient un +337 -> priorité au Fixe
        creer_ligne_test(
            {
                "Individu.Nom": "DURAND",
                "Individu.Prenom": "ALICE",
                "Structure.Nom": "1ERE FARFADETS DE LA ROCHELLE",
                "Fonction.Code": "170",
                "Fonction.CategorieMembre": "0",
                "Inscriptions.Type": "Inscrit(e)",
                "RepresentantLegal1.Nom": "DURAND",
                "RepresentantLegal1.Prenom": "CLAIRE",
                "RepresentantLegal1.TelephonePortable1": "",
                "RepresentantLegal1.TelephonePortable2": "0711223344",
                "RepresentantLegal1.TelephoneDomicile": "0699887766",
                # Cas 2 : Portable 1 vide, Fixe est un vrai fixe (+335), Portable 2 contient un +337 -> Portable 2 déplacé
                "RepresentantLegal2.Nom": "DURAND",
                "RepresentantLegal2.Prenom": "THOMAS",
                "RepresentantLegal2.TelephonePortable1": "",
                "RepresentantLegal2.TelephonePortable2": "0755667788",
                "RepresentantLegal2.TelephoneDomicile": "0546112233",
            }
        ),
    ]

    contacts = extraire_et_regrouper_parents(COLONNES_MINIMALES_TEST, lignes)
    claire, thomas = contacts[0], contacts[1]

    assert claire.prenom == "Claire"
    assert claire.telephone_portable_principal == "+33699887766"
    assert claire.telephone_portable_secondaire == "+33711223344"
    assert claire.telephone_fixe == ""

    assert thomas.prenom == "Thomas"
    assert thomas.telephone_portable_principal == "+33755667788"
    assert thomas.telephone_portable_secondaire == ""
    assert thomas.telephone_fixe == "+33546112233"


def test_generation_classeur_excel_et_cli(tmp_path: Path) -> None:
    """Vérifie la génération complète du fichier `.xlsx` via la CLI."""
    chemin_entree = tmp_path / "ExportTest.xls"
    chemin_sortie = tmp_path / "ContactsParentsTest.xlsx"

    ecrire_html_xls_synthetique(
        chemin_entree,
        [
            creer_ligne_test(
                {
                    "Individu.Nom": "DUPONT",
                    "Individu.Prenom": "LÉA",
                    "Structure.Nom": "1ERE FARFADETS DE LA ROCHELLE",
                    "Fonction.Code": "170",
                    "Fonction.CategorieMembre": "0",
                    "Inscriptions.Type": "Inscrit(e)",
                    "RepresentantLegal1.CodeAdherent": "111222333",
                    "RepresentantLegal1.Nom": "DUPONT",
                    "RepresentantLegal1.Prenom": "JEAN-PIERRE",
                    "RepresentantLegal1.TelephonePortable1": "0612345678",
                }
            )
        ],
    )

    code_retour = principale(["--entree", str(chemin_entree), "--sortie", str(chemin_sortie)])
    assert code_retour == 0
    assert chemin_sortie.is_file()

    classeur = load_workbook(chemin_sortie, data_only=True)
    assert classeur.sheetnames == [NOM_ONGLET_ORIGINE, NOM_ONGLET_PARENTS]

    feuille_parents = classeur[NOM_ONGLET_PARENTS]
    en_tetes = [cellule.value for cellule in feuille_parents[1]]
    assert tuple(en_tetes) == EN_TETES_CONTACTS_PARENTS

    premiere_ligne = [cellule.value for cellule in feuille_parents[2]]
    assert premiere_ligne[0] == "DUPONT"
    assert premiere_ligne[1] == "Jean-Pierre"
    assert premiere_ligne[2] == "+33612345678"  # Format international sans espaces
    assert premiere_ligne[6] == "Léa DUPONT"
    assert premiere_ligne[7] == "Farfadets"
    assert premiere_ligne[9] == "111222333"


def test_integration_export_individus_reel(tmp_path: Path) -> None:
    """Vérifie l'extraction sur le fichier réel `ExportIndividus.xls` sans l'altérer."""
    chemin_export = RACINE_PROJET / "ExportIndividus.xls"
    if not chemin_export.is_file():
        return

    empreinte_avant = hashlib.sha256(chemin_export.read_bytes()).hexdigest()
    chemin_sortie = tmp_path / "ContactsParents.xlsx"

    contacts = executer_extraction(chemin_entree=chemin_export, chemin_sortie=chemin_sortie)

    empreinte_apres = hashlib.sha256(chemin_export.read_bytes()).hexdigest()
    assert empreinte_avant == empreinte_apres, "`ExportIndividus.xls` ne doit subir aucune modification."

    # 1. Nombre exact de parents uniques dédoublonnés sur (Nom, Prénom) : 134
    assert len(contacts) == 134
    cles = [(c.nom, c.prenom) for c in contacts]
    assert len(cles) == len(set(cles)), "Aucun doublon (Nom, Prénom) ne doit subsister."

    # 2. Vérification du formatage NOM (MAJUSCULES), Prénom (casse titre) et téléphones (+33 sans espaces)
    for contact in contacts:
        assert contact.nom == contact.nom.upper()
        assert contact.prenom != contact.prenom.upper()
        assert contact.noms_enfants != ""
        assert contact.unites_enfants != ""
        for numero in (
            contact.telephone_portable_principal,
            contact.telephone_portable_secondaire,
            contact.telephone_fixe,
        ):
            if numero:
                assert numero.startswith("+33")
                assert " " not in numero

    # 3. Vérification de l'enrichissement des Parents Animateurs Farfadets (18 fonctions principales + 2 délégations = 20)
    parents_farfadets = [
        c for c in contacts if "Parent Animateur Farfadet" in c.remarques
    ]
    assert len(parents_farfadets) == 20
    assert sum(1 for c in parents_farfadets if c.numero_adherent != "") >= 13

    # 4. Vérification du classeur Excel généré
    classeur = load_workbook(chemin_sortie, data_only=True)
    assert classeur.sheetnames == [NOM_ONGLET_ORIGINE, NOM_ONGLET_PARENTS]
    assert classeur[NOM_ONGLET_ORIGINE].max_row == 154  # 1 en-tête + 153 lignes
    assert classeur[NOM_ONGLET_PARENTS].max_row == 135  # 1 en-tête + 134 parents
