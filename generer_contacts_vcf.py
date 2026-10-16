#!/usr/bin/env python3
"""Génération du fichier vCard (`ContactParents.vcf`) pour l'importation dans Google Contacts.

Ce script lit en lecture seule l'onglet `ContactsParents` du classeur Excel
`ContactsParents.xlsx` et génère le fichier `ContactParents.vcf` au format standard
vCard 3.0 (UTF-8) directement importable dans Google Contacts. Chaque fiche contact
contient l'identité du parent (`Prénom NOM`), ses numéros de téléphone (portable
principal, portable secondaire, fixe domicile), son adresse email, son adresse
physique, une note structurée regroupant la liste des enfants avec leurs unités,
le numéro d'adhérent (s'il existe) et d'éventuelles remarques, ainsi que les
étiquettes `"SGDF"` et `"La Rochelle"`.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

FICHIER_ENTREE_DEFAUT = "ContactsParents.xlsx"
FICHIER_SORTIE_DEFAUT = "ContactParents.vcf"

NOM_ONGLET_PARENTS = "ContactsParents"
ETIQUETTES_DEFAUT: tuple[str, ...] = ("SGDF", "La Rochelle")

EN_TETES_CONTACTS_PARENTS: tuple[str, ...] = (
    "Nom",
    "Prénom",
    "Téléphone portable principal",
    "Téléphone portable secondaire",
    "Téléphone fixe (domicile)",
    "Adresse email",
    "Noms des enfants",
    "Unités des enfants",
    "Adresse physique",
    "Numéro d'adhérent",
    "Remarques",
)


class ErreurGenerationVcf(Exception):
    """Erreur métier levée lors de la lecture du classeur Excel ou de la génération vCard."""


@dataclass(frozen=True)
class FicheContactParent:
    """Représente la fiche d'un parent prête à être exportée au format vCard 3.0."""

    nom: str
    prenom: str
    telephone_portable_principal: str = ""
    telephone_portable_secondaire: str = ""
    telephone_fixe: str = ""
    adresse_email: str = ""
    noms_enfants: str = ""
    unites_enfants: str = ""
    adresse_physique: str = ""
    numero_adherent: str = ""
    remarques: str = ""
    etiquettes: tuple[str, ...] = ETIQUETTES_DEFAUT

    @property
    def nom_complet_affiche(self) -> str:
        """Retourne le nom complet affiché dans Google Contacts sous la forme 'Prénom NOM'."""
        return " ".join(part for part in (self.prenom, self.nom) if part)


def nettoyer_texte(valeur: object) -> str:
    """Convertit une valeur de cellule Excel en chaîne nettoyée de ses espaces superflus."""
    if valeur is None:
        return ""
    return re.sub(r"\s+", " ", str(valeur).strip())


def echapper_valeur_vcard(texte: str) -> str:
    """Échappe les caractères spéciaux (RFC 2426 : `\\`, `;`, `,`, `\\n`) d'une valeur vCard 3.0."""
    texte_normalise = texte.replace("\r\n", "\n").replace("\r", "\n")
    texte_echappe = (
        texte_normalise.replace("\\", "\\\\")
        .replace(";", r"\;")
        .replace(",", r"\,")
        .replace("\n", r"\n")
    )
    return texte_echappe


def decomposer_adresse_postale(adresse_physique: str) -> tuple[str, str, str, str]:
    """Décompose une adresse postale en `(rue, ville, code_postal, pays)` pour le champ `ADR`."""
    adresse_propre = nettoyer_texte(adresse_physique)
    if not adresse_propre:
        return "", "", "", ""

    segments = [seg.strip() for seg in adresse_propre.split(",") if seg.strip()]
    motif_cp_ville = re.compile(r"^(\d{4,5})\s+(.+)$")

    if len(segments) >= 2:
        correspondance_fin = motif_cp_ville.match(segments[-1])
        if correspondance_fin:
            rue = ", ".join(segments[:-1])
            code_postal = correspondance_fin.group(1)
            ville = correspondance_fin.group(2)
            return rue, ville, code_postal, ""

        if len(segments) >= 3:
            correspondance_avant_fin = motif_cp_ville.match(segments[-2])
            if correspondance_avant_fin:
                rue = ", ".join(segments[:-2])
                code_postal = correspondance_avant_fin.group(1)
                ville = correspondance_avant_fin.group(2)
                pays = segments[-1]
                return rue, ville, code_postal, pays

    if len(segments) == 1:
        correspondance_unique = motif_cp_ville.match(segments[0])
        if correspondance_unique:
            return "", correspondance_unique.group(2), correspondance_unique.group(1), ""

    return adresse_propre, "", "", ""


def formater_enfants_et_unites(noms_enfants: str, unites_enfants: str) -> str:
    """Associe chaque enfant à son unité respective sous la forme `Prénom NOM (Unité)`."""
    enfants = [e.strip() for e in noms_enfants.split(",") if e.strip()]
    if not enfants:
        return ""

    unites = [u.strip() for u in unites_enfants.split(",") if u.strip()]
    elements_formates: list[str] = []

    for indice, enfant in enumerate(enfants):
        if indice < len(unites) and unites[indice]:
            elements_formates.append(f"{enfant} ({unites[indice]})")
        elif len(unites) == 1 and unites[0]:
            elements_formates.append(f"{enfant} ({unites[0]})")
        else:
            elements_formates.append(enfant)

    return ", ".join(elements_formates)


def construire_note_parent(fiche: FicheContactParent) -> str:
    """Construit le texte multi-lignes des informations complémentaires (`NOTE`) d'un parent."""
    lignes_note: list[str] = []

    enfants_avec_unites = formater_enfants_et_unites(
        fiche.noms_enfants,
        fiche.unites_enfants,
    )
    if enfants_avec_unites:
        lignes_note.append(f"Enfants : {enfants_avec_unites}")

    if fiche.numero_adherent:
        lignes_note.append(f"Numéro d'adhérent : {fiche.numero_adherent}")

    if fiche.remarques:
        lignes_note.append(f"Remarques : {fiche.remarques}")

    return "\n".join(lignes_note)


def formater_vcard_parent(fiche: FicheContactParent) -> str:
    """Sérialise une `FicheContactParent` en bloc `VCARD` 3.0 conforme à la RFC 2426."""
    lignes_vcard: list[str] = [
        "BEGIN:VCARD",
        "VERSION:3.0",
        f"N:{echapper_valeur_vcard(fiche.nom)};{echapper_valeur_vcard(fiche.prenom)};;;",
        f"FN:{echapper_valeur_vcard(fiche.nom_complet_affiche)}",
    ]

    if fiche.telephone_portable_principal:
        lignes_vcard.append(
            f"TEL;TYPE=CELL,VOICE,PREF:{fiche.telephone_portable_principal}"
        )
    if fiche.telephone_portable_secondaire:
        lignes_vcard.append(
            f"TEL;TYPE=CELL,VOICE:{fiche.telephone_portable_secondaire}"
        )
    if fiche.telephone_fixe:
        lignes_vcard.append(f"TEL;TYPE=HOME,VOICE:{fiche.telephone_fixe}")

    if fiche.adresse_email:
        lignes_vcard.append(f"EMAIL;TYPE=INTERNET,HOME:{fiche.adresse_email}")

    if fiche.adresse_physique:
        rue, ville, code_postal, pays = decomposer_adresse_postale(
            fiche.adresse_physique
        )
        lignes_vcard.append(
            "ADR;TYPE=HOME:;;"
            f"{echapper_valeur_vcard(rue)};"
            f"{echapper_valeur_vcard(ville)};;"
            f"{echapper_valeur_vcard(code_postal)};"
            f"{echapper_valeur_vcard(pays)}"
        )

    note_parent = construire_note_parent(fiche)
    if note_parent:
        lignes_vcard.append(f"NOTE:{echapper_valeur_vcard(note_parent)}")

    if fiche.etiquettes:
        valeurs_categories = ",".join(
            echapper_valeur_vcard(etiquette) for etiquette in fiche.etiquettes
        )
        lignes_vcard.append(f"CATEGORIES:{valeurs_categories}")

    lignes_vcard.append("END:VCARD")
    return "\r\n".join(lignes_vcard)


def lire_contacts_parents_excel(chemin_entree: Path) -> list[FicheContactParent]:
    """Lit en lecture seule l'onglet `ContactsParents` du classeur Excel `ContactsParents.xlsx`."""
    if not chemin_entree.is_file():
        raise ErreurGenerationVcf(
            f"Le fichier d'entrée '{chemin_entree}' est introuvable. "
            "Veuillez d'abord générer ContactsParents.xlsx."
        )

    try:
        classeur = load_workbook(chemin_entree, read_only=True, data_only=True)
    except (OSError, InvalidFileException, BadZipFile, KeyError, ValueError) as erreur:
        raise ErreurGenerationVcf(
            f"Impossible d'ouvrir le classeur Excel '{chemin_entree}' : {erreur}"
        ) from erreur

    try:
        if NOM_ONGLET_PARENTS not in classeur.sheetnames:
            raise ErreurGenerationVcf(
                f"L'onglet '{NOM_ONGLET_PARENTS}' est absent du classeur '{chemin_entree}'."
            )

        feuille = classeur[NOM_ONGLET_PARENTS]
        iterateur_lignes = feuille.iter_rows(values_only=True)
        premiere_ligne = next(iterateur_lignes, None)
        if premiere_ligne is None:
            raise ErreurGenerationVcf(
                f"L'onglet '{NOM_ONGLET_PARENTS}' de '{chemin_entree}' est vide."
            )

        en_tetes_lus = tuple(
            nettoyer_texte(cellule)
            for cellule in premiere_ligne[: len(EN_TETES_CONTACTS_PARENTS)]
        )
        if en_tetes_lus != EN_TETES_CONTACTS_PARENTS:
            raise ErreurGenerationVcf(
                f"En-têtes invalides dans l'onglet '{NOM_ONGLET_PARENTS}' de '{chemin_entree}'. "
                f"Attendus : {list(EN_TETES_CONTACTS_PARENTS)}."
            )

        fiches: list[FicheContactParent] = []
        nombre_colonnes = len(EN_TETES_CONTACTS_PARENTS)
        for ligne in iterateur_lignes:
            valeurs = [nettoyer_texte(cellule) for cellule in ligne[:nombre_colonnes]]
            if len(valeurs) < nombre_colonnes:
                valeurs.extend([""] * (nombre_colonnes - len(valeurs)))

            nom, prenom = valeurs[0], valeurs[1]
            if not nom and not prenom:
                continue

            fiches.append(
                FicheContactParent(
                    nom=nom,
                    prenom=prenom,
                    telephone_portable_principal=valeurs[2],
                    telephone_portable_secondaire=valeurs[3],
                    telephone_fixe=valeurs[4],
                    adresse_email=valeurs[5],
                    noms_enfants=valeurs[6],
                    unites_enfants=valeurs[7],
                    adresse_physique=valeurs[8],
                    numero_adherent=valeurs[9],
                    remarques=valeurs[10],
                )
            )

        return fiches
    finally:
        classeur.close()


def ecrire_fichier_vcf(
    chemin_sortie: Path,
    fiches: list[FicheContactParent],
) -> None:
    """Écrit l'ensemble des fiches parents au format vCard 3.0 UTF-8 dans `chemin_sortie`."""
    chemin_sortie.parent.mkdir(parents=True, exist_ok=True)
    blocs_vcard = [formater_vcard_parent(fiche) for fiche in fiches]
    contenu_final = "\r\n".join(blocs_vcard) + ("\r\n" if blocs_vcard else "")
    with chemin_sortie.open("w", encoding="utf-8", newline="") as fichier_vcf:
        fichier_vcf.write(contenu_final)


def executer_generation_vcf(
    chemin_entree: Path,
    chemin_sortie: Path,
) -> list[FicheContactParent]:
    """Exécute la chaîne complète : lecture seule de `ContactsParents.xlsx` et génération `.vcf`."""
    fiches = lire_contacts_parents_excel(chemin_entree)
    ecrire_fichier_vcf(chemin_sortie=chemin_sortie, fiches=fiches)
    return fiches


def analyser_arguments(arguments: list[str] | None = None) -> argparse.Namespace:
    """Analyse les arguments en ligne de commande (intégralement en français)."""
    analyseur = argparse.ArgumentParser(
        prog="generer_contacts_vcf.py",
        description=(
            "Génère le fichier 'ContactParents.vcf' importable dans Google Contacts "
            "à partir de l'onglet 'ContactsParents' de 'ContactsParents.xlsx'."
        ),
    )
    analyseur.add_argument(
        "-e",
        "--entree",
        type=Path,
        default=Path(FICHIER_ENTREE_DEFAUT),
        help=(
            "Chemin du classeur Excel contenant l'onglet 'ContactsParents' en entrée "
            f"(défaut : {FICHIER_ENTREE_DEFAUT})."
        ),
    )
    analyseur.add_argument(
        "-s",
        "--sortie",
        type=Path,
        default=Path(FICHIER_SORTIE_DEFAUT),
        help=(
            "Chemin du fichier vCard (.vcf) généré en sortie "
            f"(défaut : {FICHIER_SORTIE_DEFAUT})."
        ),
    )
    return analyseur.parse_args(arguments)


def principale(arguments: list[str] | None = None) -> int:
    """Point d'entrée principal du script en ligne de commande."""
    options = analyser_arguments(arguments)
    try:
        fiches = executer_generation_vcf(
            chemin_entree=options.entree,
            chemin_sortie=options.sortie,
        )
    except ErreurGenerationVcf as erreur:
        print(f"Erreur : {erreur}", file=sys.stderr)
        return 1

    print(
        f"Génération terminée avec succès : {len(fiches)} contacts parents écrits "
        f"dans le fichier '{options.sortie}'."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(principale())
