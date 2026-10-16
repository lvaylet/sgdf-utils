#!/usr/bin/env python3
"""Extraction et regroupement des contacts des parents des jeunes scouts SGDF.

Ce script lit en lecture seule le fichier d'export SGDF (`ExportIndividus.xls`),
extrait l'ensemble des représentants légaux (`1`, `2`, `3`) des jeunes du groupe,
les dédoublonne sur le couple `(Nom, Prénom)` afin de regrouper les fratries et
leurs unités courtes, enrichit les fiches avec les numéros d'adhérent et les rôles
adultes éventuels (ex. Parent Animateur Farfadet), trie les entrées par ordre
alphabétique français sur `(Nom, Prénom)` et génère le classeur Excel
`ContactsParents.xlsx` contenant les onglets `ExportIndividus` et `ContactsParents`.
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

DEFAULT_INPUT_FILE = "ExportIndividus.xls"
DEFAULT_OUTPUT_FILE = "ContactsParents.xlsx"

SOURCE_SHEET_NAME = "ExportIndividus"
PARENTS_SHEET_NAME = "ContactsParents"

PARENT_CONTACTS_HEADERS: tuple[str, ...] = (
    "Nom",
    "Prénom",
    "Téléphone portable principal",
    "Téléphone portable secondaire",
    "Téléphone fixe (domicile)",
    "Adresse email",
    "Noms des enfants",
    "Unités des enfants",
    "Rue",
    "Code postal",
    "Ville",
    "Numéro d'adhérent",
    "Remarques",
)

ADULT_ROLE_LABELS: dict[tuple[str, str], str] = {
    ("271", "M."): "Parent Animateur Farfadet",
    ("271", "Mme"): "Parent Animateur Farfadet",
    ("270", "M."): "Responsable Farfadet",
    ("270", "Mme"): "Responsable Farfadet",
    ("300", "M."): "Responsable de groupe",
    ("300", "Mme"): "Responsable de groupe",
    ("301", "M."): "Responsable de groupe adjoint",
    ("301", "Mme"): "Responsable de groupe adjointe",
    ("307", "M."): "Secrétaire de groupe",
    ("307", "Mme"): "Secrétaire de groupe",
    ("309", "M."): "Trésorier de groupe",
    ("309", "Mme"): "Trésorière de groupe",
    ("316", "M."): "Cléophas de groupe",
    ("316", "Mme"): "Cléophas de groupe",
    ("330", "M."): "Chargé de mission du groupe",
    ("330", "Mme"): "Chargée de mission du groupe",
    ("181", "M."): "Membre du réseau Impeesa",
    ("181", "Mme"): "Membre du réseau Impeesa",
    ("380", "M."): "Coordonnateur du réseau Impeesa local",
    ("380", "Mme"): "Coordonnatrice du réseau Impeesa local",
    ("210", "M."): "Responsable d'unité Louveteaux-Jeannettes",
    ("210", "Mme"): "Responsable d'unité Louveteaux-Jeannettes",
    ("213", "M."): "Chef Louveteaux-Jeannettes",
    ("213", "Mme"): "Cheftaine Louveteaux-Jeannettes",
    ("220", "M."): "Responsable d'unité Scouts-Guides",
    ("220", "Mme"): "Responsable d'unité Scouts-Guides",
    ("223", "M."): "Chef Scouts-Guides",
    ("223", "Mme"): "Cheftaine Scouts-Guides",
    ("230", "M."): "Responsable d'unité Pionniers-Caravelles",
    ("230", "Mme"): "Responsable d'unité Pionniers-Caravelles",
    ("233", "M."): "Chef Pionniers-Caravelles",
    ("233", "Mme"): "Cheftaine Pionniers-Caravelles",
}


class ContactExtractionError(Exception):
    """Erreur métier levée lors de la lecture ou de la validation de l'export SGDF."""


@dataclass(frozen=True)
class AttachedChild:
    """Représente un jeune scout rattaché à un parent."""

    first_name: str
    last_name: str
    short_unit: str
    registration_status: str

    def format_full_name(self) -> str:
        """Retourne le nom complet de l'enfant sous la forme 'Prénom NOM'."""
        return " ".join(part for part in (self.first_name, self.last_name) if part)


@dataclass
class AdultGroupProfile:
    """Représente une inscription adulte dans le groupe SGDF pour enrichir un parent."""

    last_name: str
    first_name: str
    member_id: str = ""
    roles_and_delegations: list[str] = field(default_factory=list)
    mobile_phone_1: str = ""
    mobile_phone_2: str = ""
    home_phone: str = ""
    email: str = ""
    street: str = ""
    postal_code: str = ""
    city: str = ""


@dataclass
class ParentContact:
    """Représente la fiche consolidée et dédoublonnée d'un parent de jeune(s) scout(s)."""

    last_name: str
    first_name: str
    primary_mobile_phone: str = ""
    secondary_mobile_phone: str = ""
    landline_phone: str = ""
    email_address: str = ""
    children: list[AttachedChild] = field(default_factory=list)
    units: list[str] = field(default_factory=list)
    street: str = ""
    postal_code: str = ""
    city: str = ""
    member_id: str = ""
    notes_list: list[str] = field(default_factory=list)

    def add_child(self, child: AttachedChild) -> None:
        """Ajoute un enfant et son unité sans doublon à la fiche du parent."""
        if child not in self.children:
            self.children.append(child)
        if child.short_unit and child.short_unit not in self.units:
            self.units.append(child.short_unit)
        if child.registration_status and child.registration_status != "Inscrit(e)":
            status_note = f"Enfant {child.first_name} : {child.registration_status}"
            self.add_note(status_note)

    def add_note(self, note: str) -> None:
        """Ajoute une remarque non vide sans créer de doublon."""
        cleaned_note = clean_whitespace(note)
        if cleaned_note and cleaned_note not in self.notes_list:
            self.notes_list.append(cleaned_note)

    def merge_contact_info(
        self,
        *,
        primary_mobile: str = "",
        secondary_mobile: str = "",
        landline: str = "",
        email: str = "",
        street: str = "",
        postal_code: str = "",
        city: str = "",
        member_id: str = "",
    ) -> None:
        """Conserve la première valeur non vide pour chaque coordonnée du parent."""
        if not self.primary_mobile_phone and primary_mobile:
            self.primary_mobile_phone = primary_mobile
        elif (
            not self.secondary_mobile_phone
            and primary_mobile
            and primary_mobile != self.primary_mobile_phone
        ):
            self.secondary_mobile_phone = primary_mobile

        if (
            not self.secondary_mobile_phone
            and secondary_mobile
            and secondary_mobile != self.primary_mobile_phone
        ):
            self.secondary_mobile_phone = secondary_mobile

        if not self.landline_phone and landline:
            self.landline_phone = landline
        if not self.email_address and email:
            self.email_address = email
        if not (self.street or self.postal_code or self.city) and (
            street or postal_code or city
        ):
            self.street = street
            self.postal_code = postal_code
            self.city = city
        if not self.member_id and member_id:
            self.member_id = member_id

        self.reorganize_phones()

    def reorganize_phones(self) -> None:
        """Déplace un portable valide (+336/+337) vers portable 1 s'il est vide (priorité au fixe)."""
        if self.primary_mobile_phone:
            return

        if is_valid_mobile_phone(self.landline_phone):
            self.primary_mobile_phone = self.landline_phone
            self.landline_phone = ""
            if self.secondary_mobile_phone == self.primary_mobile_phone:
                self.secondary_mobile_phone = ""
        elif is_valid_mobile_phone(self.secondary_mobile_phone):
            self.primary_mobile_phone = self.secondary_mobile_phone
            self.secondary_mobile_phone = ""

    @property
    def children_names(self) -> str:
        """Retourne la liste dédoublonnée des enfants ('Prénom NOM') séparée par ', '."""
        labels: list[str] = []
        for child in self.children:
            label = child.format_full_name()
            if label and label not in labels:
                labels.append(label)
        return ", ".join(labels)

    @property
    def children_units(self) -> str:
        """Retourne les unités courtes des enfants dans le même ordre que les enfants eux-mêmes."""
        return ", ".join(
            child.short_unit for child in self.children if child.short_unit
        )

    @property
    def notes(self) -> str:
        """Retourne les remarques consolidées séparées par ' ; '."""
        return " ; ".join(self.notes_list)

    def to_excel_row(self) -> list[str]:
        """Convertit la fiche du parent en liste des 13 valeurs de colonnes attendues."""
        return [
            self.last_name,
            self.first_name,
            self.primary_mobile_phone,
            self.secondary_mobile_phone,
            self.landline_phone,
            self.email_address,
            self.children_names,
            self.children_units,
            self.street,
            self.postal_code,
            self.city,
            self.member_id,
            self.notes,
        ]


def clean_whitespace(value: str) -> str:
    """Supprime les espaces superflus en début, fin et milieu de chaîne."""
    return re.sub(r"\s+", " ", value.strip())


def normalize_last_name(raw_last_name: str) -> str:
    """Normalise un nom de famille intégralement en MAJUSCULES."""
    return clean_whitespace(raw_last_name).upper()


def normalize_first_name(raw_first_name: str) -> str:
    """Normalise un prénom avec la première lettre de chaque partie en majuscule et le reste en minuscules."""
    text = clean_whitespace(raw_first_name).lower()
    if not text:
        return ""
    return re.sub(
        r"(^|[\s\-'])([^\s\-'])",
        lambda match: match.group(1) + match.group(2).upper(),
        text,
    )


def normalize_phone(raw_phone: str) -> str:
    """Normalise un numéro de téléphone au format international compact sans espaces (ex. +33687966975)."""
    text = raw_phone.strip()
    if not text:
        return ""

    # Supprimer l'indication facultative (0) parfois présente après l'indicatif (+33(0)6...)
    text = re.sub(r"\(\s*0\s*\)", "", text)

    has_plus_prefix = text.startswith("+")
    digits = re.sub(r"\D", "", text)
    if not digits:
        return ""

    if has_plus_prefix:
        normalized_phone = f"+{digits}"
    elif digits.startswith("00") and len(digits) > 2:
        normalized_phone = f"+{digits[2:]}"
    elif digits.startswith("0") and len(digits) > 1:
        normalized_phone = f"+33{digits[1:]}"
    elif len(digits) == 9 and digits[0] in "123456789":
        normalized_phone = f"+33{digits}"
    elif digits.startswith("33") and len(digits) == 11:
        normalized_phone = f"+{digits}"
    else:
        normalized_phone = f"+{digits}"

    # Corriger le cas +3306... (13 caractères avec un 0 superflu après +33)
    if normalized_phone.startswith("+330") and len(normalized_phone) == 13:
        normalized_phone = f"+33{normalized_phone[4:]}"

    return normalized_phone


def is_valid_mobile_phone(phone: str) -> bool:
    """Indique si le numéro normalisé est un téléphone portable français valide (+336 ou +337)."""
    return phone.startswith(("+336", "+337"))


def strip_accents(text: str) -> str:
    """Retourne une chaîne dépourvue de marques diacritiques pour le tri alphabétique français."""
    decomposed = unicodedata.normalize("NFD", text)
    return "".join(char for char in decomposed if unicodedata.category(char) != "Mn")


def parent_sort_key(contact: ParentContact) -> tuple[str, str, str, str]:
    """Clé de tri alphabétique français sur le Nom puis sur le Prénom."""
    return (
        strip_accents(contact.last_name).upper(),
        contact.last_name,
        strip_accents(contact.first_name).lower(),
        contact.first_name,
    )


def normalize_short_unit(structure_name: str, role_code: str = "") -> str:
    """Convertit un nom de structure ou code fonction SGDF en nom d'unité court normalisé."""
    code = clean_whitespace(role_code)
    upper_structure = strip_accents(clean_whitespace(structure_name)).upper()

    if code in ("170", "270", "271") or "FARFADET" in upper_structure:
        return "Farfadets"
    if (
        code in ("110", "210", "213")
        or "LOUVETEAU" in upper_structure
        or "JEANNETTE" in upper_structure
    ):
        return "Louveteaux-Jeannettes"
    if (
        code in ("120", "220", "223")
        or "SCOUTS GUIDES" in upper_structure
        or "SCOUT" in upper_structure
    ):
        return "Scouts-Guides"
    if (
        code in ("130", "230", "233")
        or "PIONNIER" in upper_structure
        or "CARAVELLE" in upper_structure
    ):
        return "Pionniers-Caravelles"
    if code in ("140", "240", "243") or "COMPAGNON" in upper_structure:
        return "Compagnons"
    return clean_whitespace(structure_name)


def normalize_adult_role(
    title: str,
    role_code: str,
    male_label: str,
    female_label: str,
) -> str:
    """Retourne le libellé lisible en français du rôle adulte au sein du groupe SGDF."""
    code = clean_whitespace(role_code)
    normalized_title = clean_whitespace(title)
    key = (code, "Mme" if normalized_title == "Mme" else "M.")
    if key in ADULT_ROLE_LABELS:
        return ADULT_ROLE_LABELS[key]

    raw_label = (
        female_label if (normalized_title == "Mme" and female_label.strip()) else male_label
    )
    text = clean_whitespace(raw_label)
    if not text:
        return ""
    return text[0].upper() + text[1:].lower()


def extract_delegation_roles(delegations_text: str, title: str = "M.") -> list[str]:
    """Extrait les rôles adultes mentionnés dans la colonne `Inscription.Delegations`."""
    text = clean_whitespace(delegations_text)
    if not text:
        return []

    roles: list[str] = []
    for segment in text.split(";"):
        chunk = segment.strip()
        if not chunk:
            continue
        match = re.match(r"^(\d+)\s*-\s*([^(]+)", chunk)
        if match:
            delegation_code = match.group(1).strip()
            raw_label = match.group(2).strip()
            role = normalize_adult_role(title, delegation_code, raw_label, raw_label)
            if role and role not in roles:
                roles.append(role)
        else:
            roles.append(chunk)
    return roles


def normalize_address_part(raw_text: str) -> str:
    """Normalise une composante d'adresse en minuscules avec la première lettre de chaque mot en majuscule."""
    text = clean_whitespace(raw_text).lower()
    if not text:
        return ""
    return re.sub(
        r"(^|[\s\-'/,])([^\s\-'/,])",
        lambda match: match.group(1) + match.group(2).upper(),
        text,
    )


def join_street_lines(
    line1: str,
    line2: str = "",
    line3: str = "",
) -> str:
    """Regroupe et normalise en casse titre les lignes de rue non vides (`Ligne1`, `Ligne2`, `Ligne3`)."""
    parts: list[str] = []
    for line in (line1, line2, line3):
        cleaned = normalize_address_part(line)
        if cleaned:
            parts.append(cleaned)
    return ", ".join(parts)


def read_individuals_export_table(input_path: Path) -> tuple[list[str], list[list[str]]]:
    """Lit en lecture seule le fichier HTML/XLS `ExportIndividus.xls` et retourne `(headers, rows)`."""
    if not input_path.is_file():
        raise ContactExtractionError(
            f"Le fichier d'entrée '{input_path}' est introuvable."
        )

    try:
        raw_bytes = input_path.read_bytes()
    except OSError as exc:
        raise ContactExtractionError(
            f"Impossible de lire le fichier '{input_path}' : {exc}"
        ) from exc

    document = BeautifulSoup(raw_bytes, "lxml")
    tr_elements = document.find_all("tr")
    if not tr_elements:
        raise ContactExtractionError(
            f"Aucun tableau HTML (`<tr>`) trouvé dans le fichier '{input_path}'."
        )

    headers = [
        clean_whitespace(cell.get_text())
        for cell in tr_elements[0].find_all(["th", "td"])
    ]
    if not headers or "Individu.Nom" not in headers:
        raise ContactExtractionError(
            f"En-têtes SGDF attendus (`Individu.Nom`) absents dans '{input_path}'."
        )

    data_rows: list[list[str]] = []
    column_count = len(headers)
    for tr_element in tr_elements[1:]:
        cells = [
            clean_whitespace(cell.get_text())
            for cell in tr_element.find_all("td")
        ]
        if not any(cells):
            continue
        if len(cells) < column_count:
            cells.extend([""] * (column_count - len(cells)))
        data_rows.append(cells[:column_count])

    return headers, data_rows


def build_adult_profiles_index(
    headers: list[str],
    rows: list[list[str]],
) -> dict[tuple[str, str], AdultGroupProfile]:
    """Construit un index des inscrits adultes (`Fonction.CategorieMembre != '0'`) par `(Nom, Prénom)`."""
    column_index = {col_name: idx for idx, col_name in enumerate(headers)}

    def get_cell(row: list[str], column: str) -> str:
        idx = column_index.get(column)
        return row[idx] if idx is not None and idx < len(row) else ""

    adult_profiles: dict[tuple[str, str], AdultGroupProfile] = {}
    for row in rows:
        category = get_cell(row, "Fonction.CategorieMembre")
        if category == "0":
            continue

        last_name = normalize_last_name(get_cell(row, "Individu.Nom"))
        first_name = normalize_first_name(get_cell(row, "Individu.Prenom"))
        if not last_name and not first_name:
            continue

        key = (last_name, first_name)
        profile = adult_profiles.get(key)
        if profile is None:
            profile = AdultGroupProfile(last_name=last_name, first_name=first_name)
            adult_profiles[key] = profile

        member_code = get_cell(row, "Individu.CodeAdherent")
        if not profile.member_id and member_code:
            profile.member_id = member_code

        title = get_cell(row, "IndividuCivilite.NomCourt")
        primary_role = normalize_adult_role(
            title=title,
            role_code=get_cell(row, "Fonction.Code"),
            male_label=get_cell(row, "Fonction.NomMasculin"),
            female_label=get_cell(row, "Fonction.NomFeminin"),
        )
        if primary_role and primary_role not in profile.roles_and_delegations:
            profile.roles_and_delegations.append(primary_role)

        for delegation_role in extract_delegation_roles(
            get_cell(row, "Inscription.Delegations"),
            title,
        ):
            if delegation_role not in profile.roles_and_delegations:
                profile.roles_and_delegations.append(delegation_role)

        if not profile.mobile_phone_1:
            profile.mobile_phone_1 = normalize_phone(
                get_cell(row, "Individu.TelephonePortable1")
            )
        if not profile.mobile_phone_2:
            profile.mobile_phone_2 = normalize_phone(
                get_cell(row, "Individu.TelephonePortable2")
            )
        if not profile.home_phone:
            profile.home_phone = normalize_phone(
                get_cell(row, "Individu.TelephoneDomicile")
            )
        if not profile.email:
            profile.email = (
                get_cell(row, "Individu.CourrielPersonnel")
                or get_cell(row, "Individu.CourrielDédiéSGDF")
            )
        adult_street = join_street_lines(
            get_cell(row, "Individu.Adresse.Ligne1"),
            get_cell(row, "Individu.Adresse.Ligne2"),
            get_cell(row, "Individu.Adresse.Ligne3"),
        )
        adult_postal_code = normalize_address_part(
            get_cell(row, "Individu.Adresse.CodePostal")
        )
        adult_city = normalize_address_part(
            get_cell(row, "Individu.Adresse.Municipalite")
        )
        if not (profile.street or profile.postal_code or profile.city) and (
            adult_street or adult_postal_code or adult_city
        ):
            profile.street = adult_street
            profile.postal_code = adult_postal_code
            profile.city = adult_city

    return adult_profiles


def extract_and_group_parents(
    headers: list[str],
    rows: list[list[str]],
) -> list[ParentContact]:
    """Extrait, dédoublonne, enrichit et trie par ordre alphabétique tous les parents de jeunes."""
    column_index = {col_name: idx for idx, col_name in enumerate(headers)}

    def get_cell(row: list[str], column: str) -> str:
        idx = column_index.get(column)
        return row[idx] if idx is not None and idx < len(row) else ""

    adult_profiles = build_adult_profiles_index(headers, rows)
    parents_by_key: dict[tuple[str, str], ParentContact] = {}

    for row in rows:
        if get_cell(row, "Fonction.CategorieMembre") != "0":
            continue

        child = AttachedChild(
            first_name=normalize_first_name(get_cell(row, "Individu.Prenom")),
            last_name=normalize_last_name(get_cell(row, "Individu.Nom")),
            short_unit=normalize_short_unit(
                get_cell(row, "Structure.Nom"),
                get_cell(row, "Fonction.Code"),
            ),
            registration_status=get_cell(row, "Inscriptions.Type") or "Inscrit(e)",
        )

        for prefix in ("RepresentantLegal1", "RepresentantLegal2", "RepresentantLegal3"):
            parent_last_name = normalize_last_name(get_cell(row, f"{prefix}.Nom"))
            parent_first_name = normalize_first_name(get_cell(row, f"{prefix}.Prenom"))
            if not parent_last_name and not parent_first_name:
                continue

            parent_key = (parent_last_name, parent_first_name)
            contact = parents_by_key.get(parent_key)
            if contact is None:
                contact = ParentContact(
                    last_name=parent_last_name,
                    first_name=parent_first_name,
                )
                parents_by_key[parent_key] = contact

            parent_email = (
                get_cell(row, f"{prefix}.CourrielPersonnel")
                or get_cell(row, f"{prefix}.CourrielDédiéSGDF")
            )
            parent_street = join_street_lines(
                get_cell(row, f"{prefix}.Adresse.Ligne1"),
                get_cell(row, f"{prefix}.Adresse.Ligne2"),
                get_cell(row, f"{prefix}.Adresse.Ligne3"),
            )
            parent_postal_code = normalize_address_part(
                get_cell(row, f"{prefix}.Adresse.CodePostal")
            )
            parent_city = normalize_address_part(
                get_cell(row, f"{prefix}.Adresse.Municipalite")
            )

            contact.merge_contact_info(
                primary_mobile=normalize_phone(
                    get_cell(row, f"{prefix}.TelephonePortable1")
                ),
                secondary_mobile=normalize_phone(
                    get_cell(row, f"{prefix}.TelephonePortable2")
                ),
                landline=normalize_phone(
                    get_cell(row, f"{prefix}.TelephoneDomicile")
                ),
                email=parent_email,
                street=parent_street,
                postal_code=parent_postal_code,
                city=parent_city,
                member_id=get_cell(row, f"{prefix}.CodeAdherent"),
            )
            contact.add_child(child)

    # Enrichissement avec l'index des profils adultes inscrits dans le groupe
    for parent_key, contact in parents_by_key.items():
        adult_profile = adult_profiles.get(parent_key)
        if adult_profile is None:
            continue

        contact.merge_contact_info(
            primary_mobile=adult_profile.mobile_phone_1,
            secondary_mobile=adult_profile.mobile_phone_2,
            landline=adult_profile.home_phone,
            email=adult_profile.email,
            street=adult_profile.street,
            postal_code=adult_profile.postal_code,
            city=adult_profile.city,
            member_id=adult_profile.member_id,
        )

        # Insérer les rôles adultes en tête des remarques (avant les statuts d'inscription des enfants)
        existing_notes = list(contact.notes_list)
        contact.notes_list.clear()
        for role in adult_profile.roles_and_delegations:
            contact.add_note(role)
        for child_note in existing_notes:
            contact.add_note(child_note)

    return sorted(parents_by_key.values(), key=parent_sort_key)


def auto_fit_column_widths(worksheet) -> None:
    """Ajuste automatiquement la largeur des colonnes d'une feuille Excel pour une lecture confortable."""
    for col_index, column in enumerate(worksheet.columns, start=1):
        max_length = 0
        for cell in column:
            text_value = str(cell.value or "")
            max_length = max(max_length, len(text_value))
        column_letter = get_column_letter(col_index)
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 3, 12), 55)


def write_parent_contacts_workbook(
    output_path: Path,
    source_headers: list[str],
    source_rows: list[list[str]],
    parent_contacts: list[ParentContact],
) -> None:
    """Génère le classeur Excel `.xlsx` contenant les onglets `ExportIndividus` et `ContactsParents`."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    header_font = Font(bold=True)
    header_alignment = Alignment(vertical="center")

    # Onglet 1 : ExportIndividus (copie fidèle des données d'origine au format texte)
    source_sheet = workbook.active
    source_sheet.title = SOURCE_SHEET_NAME
    source_sheet.append(source_headers)
    for cell in source_sheet[1]:
        cell.font = header_font
        cell.alignment = header_alignment
        cell.number_format = "@"

    for data_row in source_rows:
        source_sheet.append(data_row)
        for cell in source_sheet[source_sheet.max_row]:
            cell.number_format = "@"
    source_sheet.freeze_panes = "A2"

    # Onglet 2 : ContactsParents (liste dédoublonnée et triée des parents)
    parents_sheet = workbook.create_sheet(title=PARENTS_SHEET_NAME)
    parents_sheet.append(list(PARENT_CONTACTS_HEADERS))
    for cell in parents_sheet[1]:
        cell.font = header_font
        cell.alignment = header_alignment
        cell.number_format = "@"

    for contact in parent_contacts:
        parents_sheet.append(contact.to_excel_row())
        for cell in parents_sheet[parents_sheet.max_row]:
            cell.number_format = "@"
    parents_sheet.freeze_panes = "A2"
    auto_fit_column_widths(parents_sheet)

    workbook.save(output_path)


def run_extraction(input_path: Path, output_path: Path) -> list[ParentContact]:
    """Exécute la chaîne complète : lecture seule, extraction, dédoublonnage, tri et écriture Excel."""
    headers, rows = read_individuals_export_table(input_path)
    parent_contacts = extract_and_group_parents(headers, rows)
    write_parent_contacts_workbook(
        output_path=output_path,
        source_headers=headers,
        source_rows=rows,
        parent_contacts=parent_contacts,
    )
    return parent_contacts


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    """Analyse les arguments en ligne de commande (descriptions d'aide en français)."""
    parser = argparse.ArgumentParser(
        prog="extract_parent_contacts.py",
        description=(
            "Extrait, dédoublonne et trie les contacts des parents des jeunes scouts SGDF "
            "depuis l'export intranet vers l'onglet 'ContactsParents'."
        ),
    )
    parser.add_argument(
        "-i",
        "--input",
        type=Path,
        default=Path(DEFAULT_INPUT_FILE),
        help=f"Chemin du fichier d'export SGDF en entrée (défaut : {DEFAULT_INPUT_FILE}).",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path(DEFAULT_OUTPUT_FILE),
        help=f"Chemin du classeur Excel généré en sortie (défaut : {DEFAULT_OUTPUT_FILE}).",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    """Point d'entrée principal du script en ligne de commande."""
    args = parse_args(argv)
    try:
        contacts = run_extraction(
            input_path=args.input,
            output_path=args.output,
        )
    except ContactExtractionError as exc:
        print(f"Erreur : {exc}", file=sys.stderr)
        return 1

    print(
        f"Extraction terminée avec succès : {len(contacts)} parents uniques enregistrés "
        f"dans l'onglet '{PARENTS_SHEET_NAME}' du fichier '{args.output}'."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
