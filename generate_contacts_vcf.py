#!/usr/bin/env python3
"""Génération du fichier vCard (`ContactParents.vcf`) pour l'importation dans Google Contacts.

Ce script lit en lecture seule l'onglet `ContactsParents` du classeur Excel
`ContactsParents.xlsx` et génère le fichier `ContactParents.vcf` au format standard
vCard 3.0 (UTF-8) directement importable dans Google Contacts. Chaque fiche contact
contient l'identité du parent (`Prénom NOM`), ses numéros de téléphone (portable
principal, portable secondaire, fixe domicile), son adresse email, son adresse
physique, une note structurée regroupant la liste des enfants avec leurs unités,
le numéro d'adhérent (s'il existe) et d'éventuelles remarques, ainsi que les
étiquettes `"SGDF"` et `"La Rochelle"`.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

DEFAULT_INPUT_FILE = "ContactsParents.xlsx"
DEFAULT_OUTPUT_FILE = "ContactParents.vcf"

PARENTS_SHEET_NAME = "ContactsParents"
DEFAULT_CATEGORIES: tuple[str, ...] = ("SGDF", "La Rochelle")

PARENT_CONTACTS_HEADERS: tuple[str, ...] = (
    "Nom",
    "Prénom",
    "Téléphone portable principal",
    "Téléphone portable secondaire",
    "Téléphone fixe (domicile)",
    "Adresse email",
    "Noms des enfants",
    "Unités des enfants",
    "Rue",
    "Code postal",
    "Ville",
    "Numéro d'adhérent",
    "Remarques",
)


class VcfGenerationError(Exception):
    """Erreur métier levée lors de la lecture du classeur Excel ou de la génération vCard."""


@dataclass(frozen=True)
class ParentContactCard:
    """Représente la fiche d'un parent prête à être exportée au format vCard 3.0."""

    last_name: str
    first_name: str
    primary_mobile_phone: str = ""
    secondary_mobile_phone: str = ""
    landline_phone: str = ""
    email_address: str = ""
    children_names: str = ""
    children_units: str = ""
    street: str = ""
    postal_code: str = ""
    city: str = ""
    member_id: str = ""
    notes: str = ""
    categories: tuple[str, ...] = DEFAULT_CATEGORIES

    @property
    def formatted_full_name(self) -> str:
        """Retourne le nom complet affiché dans Google Contacts sous la forme 'Prénom NOM'."""
        return " ".join(part for part in (self.first_name, self.last_name) if part)


def clean_text(value: object) -> str:
    """Convertit une valeur de cellule Excel en chaîne nettoyée de ses espaces superflus."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def escape_vcard_value(text: str) -> str:
    """Échappe les caractères spéciaux (RFC 2426 : `\\`, `;`, `,`, `\\n`) d'une valeur vCard 3.0."""
    normalized_text = text.replace("\r\n", "\n").replace("\r", "\n")
    escaped_text = (
        normalized_text.replace("\\", "\\\\")
        .replace(";", r"\;")
        .replace(",", r"\,")
        .replace("\n", r"\n")
    )
    return escaped_text


def format_children_with_units(children_names: str, children_units: str) -> str:
    """Associe chaque enfant à son unité respective sous la forme `Prénom NOM (Unité)`."""
    children = [child.strip() for child in children_names.split(",") if child.strip()]
    if not children:
        return ""

    units = [unit.strip() for unit in children_units.split(",") if unit.strip()]
    formatted_items: list[str] = []

    for index, child in enumerate(children):
        if index < len(units) and units[index]:
            formatted_items.append(f"{child} ({units[index]})")
        elif len(units) == 1 and units[0]:
            formatted_items.append(f"{child} ({units[0]})")
        else:
            formatted_items.append(child)

    return ", ".join(formatted_items)


def build_parent_note(card: ParentContactCard) -> str:
    """Construit le texte multi-lignes des informations complémentaires (`NOTE`) d'un parent."""
    note_lines: list[str] = []

    children_with_units = format_children_with_units(
        card.children_names,
        card.children_units,
    )
    if children_with_units:
        note_lines.append(f"Enfants : {children_with_units}")

    if card.member_id:
        note_lines.append(f"Numéro d'adhérent : {card.member_id}")

    if card.notes:
        note_lines.append(f"Remarques : {card.notes}")

    return "\n".join(note_lines)


def format_parent_vcard(card: ParentContactCard) -> str:
    """Sérialise une `ParentContactCard` en bloc `VCARD` 3.0 conforme à la RFC 2426."""
    vcard_lines: list[str] = [
        "BEGIN:VCARD",
        "VERSION:3.0",
        f"N:{escape_vcard_value(card.last_name)};{escape_vcard_value(card.first_name)};;;",
        f"FN:{escape_vcard_value(card.formatted_full_name)}",
    ]

    if card.primary_mobile_phone:
        vcard_lines.append(
            f"TEL;TYPE=CELL,VOICE,PREF:{card.primary_mobile_phone}"
        )
    if card.secondary_mobile_phone:
        vcard_lines.append(
            f"TEL;TYPE=CELL,VOICE:{card.secondary_mobile_phone}"
        )
    if card.landline_phone:
        vcard_lines.append(f"TEL;TYPE=HOME,VOICE:{card.landline_phone}")

    if card.email_address:
        vcard_lines.append(f"EMAIL;TYPE=INTERNET,HOME:{card.email_address}")

    if card.street or card.postal_code or card.city:
        vcard_lines.append(
            "ADR;TYPE=HOME:;;"
            f"{escape_vcard_value(card.street)};"
            f"{escape_vcard_value(card.city)};;"
            f"{escape_vcard_value(card.postal_code)};"
        )

    parent_note = build_parent_note(card)
    if parent_note:
        vcard_lines.append(f"NOTE:{escape_vcard_value(parent_note)}")

    if card.categories:
        category_values = ",".join(
            escape_vcard_value(category) for category in card.categories
        )
        vcard_lines.append(f"CATEGORIES:{category_values}")

    vcard_lines.append("END:VCARD")
    return "\r\n".join(vcard_lines)


def read_parent_contacts_excel(input_path: Path) -> list[ParentContactCard]:
    """Lit en lecture seule l'onglet `ContactsParents` du classeur Excel `ContactsParents.xlsx`."""
    if not input_path.is_file():
        raise VcfGenerationError(
            f"Le fichier d'entrée '{input_path}' est introuvable. "
            "Veuillez d'abord générer ContactsParents.xlsx."
        )

    try:
        workbook = load_workbook(input_path, read_only=True, data_only=True)
    except (OSError, InvalidFileException, BadZipFile, KeyError, ValueError) as exc:
        raise VcfGenerationError(
            f"Impossible d'ouvrir le classeur Excel '{input_path}' : {exc}"
        ) from exc

    try:
        if PARENTS_SHEET_NAME not in workbook.sheetnames:
            raise VcfGenerationError(
                f"L'onglet '{PARENTS_SHEET_NAME}' est absent du classeur '{input_path}'."
            )

        worksheet = workbook[PARENTS_SHEET_NAME]
        row_iterator = worksheet.iter_rows(values_only=True)
        first_row = next(row_iterator, None)
        if first_row is None:
            raise VcfGenerationError(
                f"L'onglet '{PARENTS_SHEET_NAME}' de '{input_path}' est vide."
            )

        read_headers = tuple(
            clean_text(cell)
            for cell in first_row[: len(PARENT_CONTACTS_HEADERS)]
        )
        if read_headers != PARENT_CONTACTS_HEADERS:
            raise VcfGenerationError(
                f"En-têtes invalides dans l'onglet '{PARENTS_SHEET_NAME}' de '{input_path}'. "
                f"Attendus : {list(PARENT_CONTACTS_HEADERS)}."
            )

        cards: list[ParentContactCard] = []
        column_count = len(PARENT_CONTACTS_HEADERS)
        for row in row_iterator:
            values = [clean_text(cell) for cell in row[:column_count]]
            if len(values) < column_count:
                values.extend([""] * (column_count - len(values)))

            last_name, first_name = values[0], values[1]
            if not last_name and not first_name:
                continue

            cards.append(
                ParentContactCard(
                    last_name=last_name,
                    first_name=first_name,
                    primary_mobile_phone=values[2],
                    secondary_mobile_phone=values[3],
                    landline_phone=values[4],
                    email_address=values[5],
                    children_names=values[6],
                    children_units=values[7],
                    street=values[8],
                    postal_code=values[9],
                    city=values[10],
                    member_id=values[11],
                    notes=values[12],
                )
            )

        return cards
    finally:
        workbook.close()


def write_vcf_file(
    output_path: Path,
    cards: list[ParentContactCard],
) -> None:
    """Écrit l'ensemble des fiches parents au format vCard 3.0 UTF-8 dans `output_path`."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    vcard_blocks = [format_parent_vcard(card) for card in cards]
    final_content = "\r\n".join(vcard_blocks) + ("\r\n" if vcard_blocks else "")
    with output_path.open("w", encoding="utf-8", newline="") as vcf_file:
        vcf_file.write(final_content)


def run_vcf_generation(
    input_path: Path,
    output_path: Path,
) -> list[ParentContactCard]:
    """Exécute la chaîne complète : lecture seule de `ContactsParents.xlsx` et génération `.vcf`."""
    cards = read_parent_contacts_excel(input_path)
    write_vcf_file(output_path=output_path, cards=cards)
    return cards


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    """Analyse les arguments en ligne de commande (descriptions d'aide en français)."""
    parser = argparse.ArgumentParser(
        prog="generate_contacts_vcf.py",
        description=(
            "Génère le fichier 'ContactParents.vcf' importable dans Google Contacts "
            "à partir de l'onglet 'ContactsParents' de 'ContactsParents.xlsx'."
        ),
    )
    parser.add_argument(
        "-i",
        "--input",
        type=Path,
        default=Path(DEFAULT_INPUT_FILE),
        help=(
            "Chemin du classeur Excel contenant l'onglet 'ContactsParents' en entrée "
            f"(défaut : {DEFAULT_INPUT_FILE})."
        ),
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path(DEFAULT_OUTPUT_FILE),
        help=(
            "Chemin du fichier vCard (.vcf) généré en sortie "
            f"(défaut : {DEFAULT_OUTPUT_FILE})."
        ),
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    """Point d'entrée principal du script en ligne de commande."""
    args = parse_args(argv)
    try:
        cards = run_vcf_generation(
            input_path=args.input,
            output_path=args.output,
        )
    except VcfGenerationError as exc:
        print(f"Erreur : {exc}", file=sys.stderr)
        return 1

    print(
        f"Génération terminée avec succès : {len(cards)} contacts parents écrits "
        f"dans le fichier '{args.output}'."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
