"""Génération du classeur Excel `ContactsParents.xlsx` sur disque (`Path`) ou en mémoire (`BytesIO`)."""

from __future__ import annotations

import io
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from sgdf_contacts.domain import (
    PARENT_CONTACTS_HEADERS,
    PARENTS_SHEET_NAME,
    SOURCE_SHEET_NAME,
    ParentContact,
)


def auto_fit_column_widths(worksheet) -> None:
    """Ajuste automatiquement la largeur des colonnes d'une feuille Excel pour une lecture confortable."""
    for col_index, column in enumerate(worksheet.columns, start=1):
        max_length = 0
        for cell in column:
            text_value = str(cell.value or "")
            max_length = max(max_length, len(text_value))
        column_letter = get_column_letter(col_index)
        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 3, 12), 55
        )


def write_parent_contacts_workbook(
    output_path: Path | BinaryIO | None = None,
    source_headers: list[str] | None = None,
    source_rows: list[list[str]] | None = None,
    parent_contacts: list[ParentContact] | None = None,
) -> bytes:
    """Génère le classeur Excel `.xlsx` contenant les onglets `ExportIndividus` et `ContactsParents`."""
    headers = source_headers or []
    rows = source_rows or []
    contacts = parent_contacts or []

    workbook = Workbook()
    header_font = Font(bold=True)
    header_alignment = Alignment(vertical="center")

    # Onglet 1 : ExportIndividus (copie fidèle des données d'origine au format texte)
    source_sheet = workbook.active
    source_sheet.title = SOURCE_SHEET_NAME
    source_sheet.append(headers)
    for cell in source_sheet[1]:
        cell.font = header_font
        cell.alignment = header_alignment
        cell.number_format = "@"

    for data_row in rows:
        source_sheet.append(data_row)
        for cell in source_sheet[source_sheet.max_row]:
            cell.number_format = "@"
    source_sheet.freeze_panes = "A2"

    # Onglet 2 : ContactsParents (liste dédoublonnée et triée des parents)
    parents_sheet = workbook.create_sheet(title=PARENTS_SHEET_NAME)
    parents_sheet.append(list(PARENT_CONTACTS_HEADERS))
    for cell in parents_sheet[1]:
        cell.font = header_font
        cell.alignment = header_alignment
        cell.number_format = "@"

    for contact in contacts:
        parents_sheet.append(contact.to_excel_row())
        for cell in parents_sheet[parents_sheet.max_row]:
            cell.number_format = "@"
    parents_sheet.freeze_panes = "A2"
    auto_fit_column_widths(parents_sheet)

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook_bytes = buffer.getvalue()

    if output_path is not None:
        if hasattr(output_path, "write"):
            output_path.write(workbook_bytes)
            if hasattr(output_path, "seek"):
                output_path.seek(0)
        else:
            target_path = Path(output_path)
            target_path.parent.mkdir(parents=True, exist_ok=True)
            target_path.write_bytes(workbook_bytes)

    return workbook_bytes
