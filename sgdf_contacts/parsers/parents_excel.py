"""Lecture de l'onglet `ContactsParents` depuis le classeur Excel `ContactsParents.xlsx`."""

from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from sgdf_contacts.domain import (
    PARENT_CONTACTS_HEADERS,
    PARENTS_SHEET_NAME,
    ParentContactCard,
    VcfGenerationError,
    clean_text,
)


def read_parent_contacts_excel(input_path: Path) -> list[ParentContactCard]:
    """Lit en lecture seule l'onglet `ContactsParents` du classeur Excel `ContactsParents.xlsx`."""
    if not input_path.is_file():
        raise VcfGenerationError(
            f"Le fichier d'entrée '{input_path}' est introuvable. "
            "Veuillez d'abord générer ContactsParents.xlsx."
        )

    try:
        workbook = load_workbook(input_path, read_only=True, data_only=True)
    except (OSError, InvalidFileException, BadZipFile, KeyError, ValueError) as exc:
        raise VcfGenerationError(
            f"Impossible d'ouvrir le classeur Excel '{input_path}' : {exc}"
        ) from exc

    try:
        if PARENTS_SHEET_NAME not in workbook.sheetnames:
            raise VcfGenerationError(
                f"L'onglet '{PARENTS_SHEET_NAME}' est absent du classeur '{input_path}'."
            )

        worksheet = workbook[PARENTS_SHEET_NAME]
        row_iterator = worksheet.iter_rows(values_only=True)
        first_row = next(row_iterator, None)
        if first_row is None:
            raise VcfGenerationError(
                f"L'onglet '{PARENTS_SHEET_NAME}' de '{input_path}' est vide."
            )

        read_headers = tuple(
            clean_text(cell)
            for cell in first_row[: len(PARENT_CONTACTS_HEADERS)]
        )
        if read_headers != PARENT_CONTACTS_HEADERS:
            raise VcfGenerationError(
                f"En-têtes invalides dans l'onglet '{PARENTS_SHEET_NAME}' de '{input_path}'. "
                f"Attendus : {list(PARENT_CONTACTS_HEADERS)}."
            )

        cards: list[ParentContactCard] = []
        column_count = len(PARENT_CONTACTS_HEADERS)
        for row in row_iterator:
            values = [clean_text(cell) for cell in row[:column_count]]
            if len(values) < column_count:
                values.extend([""] * (column_count - len(values)))

            last_name, first_name = values[0], values[1]
            if not last_name and not first_name:
                continue

            cards.append(
                ParentContactCard(
                    last_name=last_name,
                    first_name=first_name,
                    primary_mobile_phone=values[2],
                    secondary_mobile_phone=values[3],
                    landline_phone=values[4],
                    email_address=values[5],
                    children_names=values[6],
                    children_units=values[7],
                    street=values[8],
                    postal_code=values[9],
                    city=values[10],
                    member_id=values[11],
                    notes=values[12],
                )
            )

        return cards
    finally:
        workbook.close()
